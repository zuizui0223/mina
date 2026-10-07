#!/usr/bin/env python3
"""Build Supporting Information DOCX for the spatial-recovery Ecology Article."""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TITLE = "Population recovery can retrace local losses while spatial structure remains altered"

FIGURES = [
    ("Figure S1", "figure3_ross_path_state_calibration.png"),
    ("Figure S2", "figure4_process_contrasts.png"),
]


def _page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _set_layout(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def _set_font(run, size=12, bold=False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _format_doc(doc: Document) -> None:
    for section in doc.sections:
        _set_layout(section)
        section.footer.is_linked_to_previous = False
        _page_number(section.footer.paragraphs[0])
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for p in doc.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        for run in p.runs:
            _set_font(run)


def _captions_from_source(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    out: dict[str, str] = {}
    current = None
    buf: list[str] = []
    for line in text.splitlines():
        if line.startswith("## Figure S"):
            if current is not None:
                out[current] = "\n".join(buf).strip()
            current = line[3:].strip()
            buf = []
        elif current is not None:
            if line.startswith("## "):
                out[current] = "\n".join(buf).strip()
                current = None
                buf = []
            else:
                buf.append(line)
    if current is not None:
        out[current] = "\n".join(buf).strip()
    return out


def build(supplement: Path, figure_dir: Path, out: Path) -> None:
    captions = _captions_from_source(supplement)
    doc = Document()
    _format_doc(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supporting Information")
    _set_font(r, size=14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITLE)
    _set_font(r, size=12, bold=True)

    p = doc.add_paragraph()
    p.add_run(
        "This file contains descriptive supporting figures referenced by the "
        "main manuscript. Figure S1 is within-Ross calibration and is not an "
        "independent replication. Figure S2 contains process contrasts and is "
        "not a replication of the focal Ross disturbance/rebound natural experiment."
    )

    for label, filename in FIGURES:
        doc.add_page_break()

        p = doc.add_paragraph()
        r = p.add_run(label + ".")
        _set_font(r, bold=True)

        key = next((k for k in captions if k.startswith(label)), None)
        if key is None:
            raise ValueError(f"caption not found for {label}")
        cap = captions[key]
        source_line = "**Source image:** " + chr(96) + filename + chr(96)
        cap = cap.replace(source_line, "").strip()
        cap = cap.replace("**", "").replace(chr(96), "")
        doc.add_paragraph(cap)

        image = figure_dir / filename
        if not image.exists():
            raise FileNotFoundError(image)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(6.7))

    _format_doc(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--supplement", type=Path, required=True)
    ap.add_argument("--figure-dir", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    build(args.supplement, args.figure_dir, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
