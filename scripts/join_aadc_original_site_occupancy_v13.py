#!/usr/bin/env python3
"""V13 ORIGINAL AADC eight-seabird survey × East Antarctic potential-site join.

Use original public AADC API object files, exact V12 archived SHA256s,
published headers and source observation categories. Report sample support
ONLY: no island × nesting-guild ecological regression, no artificial zero
for unknown/nonreporting and no arbitrary resolution of conflicting surveys.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re

from openpyxl import load_workbook
import requests
import xlrd

from audit_aadc_original_workbook_headers_v11 import (
    DATASETS,HOST,SOURCE_FILES,public_dataset_manifests,trusted_response_get,MIME_LIMIT
)
from decode_aadc_original_site_occupancy_headers_v12 import (
    PINNED_FILES,ADDITIONAL_GEO_NAME
)

SOURCE_SPECS={
    "potential_geography_2021": ("4344",ADDITIONAL_GEO_NAME,
        "2236269e754962f9af082a0853993fc65a4afd20c7489a65fae81bdd1899ffaa",
        "2_Geographic sites",4884),
    "potential_geography_2025": ("5959",SOURCE_FILES["geographic_potential_sites"],
        PINNED_FILES["geographic_potential_sites"],"2_geog_sites_all",7785),
    "penguin_original_search": ("5959",SOURCE_FILES["penguin_original_occupancy"],
        PINNED_FILES["penguin_original_occupancy"],"2_search events",11089),
    "flying_seabird_original_search": ("5959",SOURCE_FILES["flying_seabird_original_occupancy"],
        PINNED_FILES["flying_seabird_original_occupancy"],"2_search events",22018),
    "survey_certainty_rules": ("5959",SOURCE_FILES["occupancy_certainty_criteria"],
        PINNED_FILES["occupancy_certainty_criteria"],"2_occurrence certainty",216),
}
REQUIRED_HEADERS={
    "potential_geography_2021":{
        "Geog_Site_ID","Geog_Site_Latitude","Geog_Site_Longitude",
        "Geog_Site_Area",
    },
    "potential_geography_2025":{
        "geog_site_id","geog_site_latitude","geog_site_longitude",
        "geog_site_feature","geog_site_area","geog_site_area_cat3",
        "sub_group","distance_station","distance_inland_coast",
    },
    "penguin_original_search":{
        "geog_site_id","occurrence","vantage","species_code",
        "species_expert_validation","site_expert_validation",
        "occurrence_expert_validation","season","obs_number",
    },
    "flying_seabird_original_search":{
        "geog_site_id","occurrence","vantage","species_code",
        "species_expert_validation","site_expert_validation",
        "occurrence_expert_validation","season","obs_number",
    },
    "survey_certainty_rules":{
        "species_code","occurrence","vantage","geog_site_area_cat3",
        "rel_sp_detect","rel_vantage_detect","site_search_completeness",
        "product","occurrence_cert_cat"
    }
}
OUT_SURVEY_COLS=(
    "source_study","geog_site_id","sub_group","species_code","occurrence",
    "vantage","day","day_uncertainty","month","year","season","obs_number",
    "breed_site_id","species_expert_validation","site_expert_validation",
    "occurrence_expert_validation","source_geog_site_lookup",
    "source_geog_site_feature","source_geog_site_area",
    "source_geog_site_area_cat3","source_latitude","source_longitude",
    "source_distance_station","source_distance_inland_coast",
    "certainty_lookup","rel_sp_detect","rel_vantage_detect",
    "site_search_completeness","certainty_product","occurrence_cert_cat",
    "scientific_absence_is_not_inferred_from_nonreporting"
)
OUT_SITE_COLS=(
    "geog_site_id","sub_group","geog_site_feature","geog_site_area",
    "geog_site_area_cat3","latitude","longitude","distance_station",
    "distance_inland_coast","was_in_2021_original_geography",
)


def source_text(value):
    if value is None: return ""
    if isinstance(value, bool): return str(value)
    if isinstance(value,float):
        if math.isnan(value): return ""
        if value.is_integer(): return str(int(value))
    return str(value).strip()


def source_id(value):
    v=source_text(value)
    if not v: return ""
    return v.upper()


def sheet_source(raw: bytes, sheet_name: str):
    """Read explicitly named original public Excel source (no formulas run)."""
    if raw.startswith(b"PK"):
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                raise ValueError("SOURCE_STOP no exact original worksheet "+sheet_name)
            sheet=wb[sheet_name]
            itr=sheet.iter_rows(values_only=True)
            header=[source_text(v) for v in next(itr)]
            cells=(list(row) for row in itr)
            rows=[]
            for entry in cells:
                if not any(source_text(v) for v in entry):
                    continue
                if any(source_text(v) for v in entry[len(header):]):
                    raise ValueError("SCHEMA_STOP values beyond declared source columns")
                obj={field:entry[i] if i<len(entry) else None
                     for i,field in enumerate(header) if field}
                rows.append(obj)
        finally:
            wb.close()
    elif raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        wb=xlrd.open_workbook(file_contents=raw,on_demand=True)
        try:
            if sheet_name not in wb.sheet_names():
                raise ValueError("SOURCE_STOP no exact legacy worksheet "+sheet_name)
            sheet=wb.sheet_by_name(sheet_name)
            header=[source_text(x) for x in sheet.row_values(0)]
            rows=[]
            for i in range(1,sheet.nrows):
                cell=sheet.row_values(i)
                if not any(source_text(v) for v in cell):continue
                if any(source_text(v) for v in cell[len(header):]):
                    raise ValueError("SCHEMA_STOP legacy source surplus columns")
                rows.append({field:cell[k] if k<len(cell) else None
                            for k,field in enumerate(header) if field})
        finally:
            wb.release_resources()
    else:
        raise ValueError("SOURCE_STOP original author spreadsheet not Excel")
    # The 2021 AADC geography XLSX has an author-populated A:I header,
    # but the source worksheet's declared extent reaches column AK.
    # Blank trailing heading cells are valid; ONLY duplicate NONBLANK
    # column names cause ambiguity. Never invent labels for blank cells.
    populated=[name for name in header if name]
    if not populated or len(populated)!=len(set(populated)):
        raise ValueError("SCHEMA_STOP duplicate original nonblank headings")
    return header,rows


def fetch_original_sources(session,manifests):
    sources={}
    source_details={}
    for label,(eds,filename,expected_hash,sheet,count) in SOURCE_SPECS.items():
        manifest=manifests[eds]["original_object_manifest"]["file_descriptors"]
        if filename not in manifest:
            raise ValueError("SOURCE_STOP published file missing: "+filename)
        uid=manifests[eds]["uuid"]
        from urllib.parse import quote
        url=f"https://{HOST}/eds/api/dataset/{quote(uid)}/object/download?prefix={quote(filename,safe='')}"
        raw,meta=trusted_response_get(session,url,MIME_LIMIT)
        if meta["source_sha256"]!=expected_hash:
            raise ValueError("SOURCE_STOP original AADC file changed: "+label)
        header,rows=sheet_source(raw,sheet)
        absent=REQUIRED_HEADERS[label]-set(header)
        if absent:raise ValueError(f"SCHEMA_STOP {label} missing {sorted(absent)}")
        if len(rows)!=count:
            raise ValueError(f"SOURCE_VERSION_STOP expected {count} original rows, got {len(rows)} in {label}")
        sources[label]=rows
        source_details[label]={
            "original_DOI":DATASETS[eds]["doi"],
            "original_filename":filename,
            "original_worksheet":sheet,
            "original_SHA256":meta["source_sha256"],
            "actual_source_file_bytes":len(raw),
            "original_observation_rows":len(rows),
            "original_columns":header,
        }
    return sources,source_details


def index_geography(rows):
    sites={}
    duplicates=Counter()
    for src in rows:
        sid=source_id(src.get("geog_site_id"))
        if not sid:
            duplicates["EMPTY_ID"]+=1
            continue
        if sid in sites:
            duplicates["DUPLICATE_ID"]+=1
            continue
        sites[sid]=src
    if duplicates:
        raise ValueError("GEOGRAPHY_STOP original site IDs missing/duplicated: "+str(dict(duplicates)))
    return sites


def build_certainty(rules):
    lookup={}
    repeated=Counter()
    empty=Counter()
    for row in rules:
        key=tuple(source_text(row.get(c)).casefold()
                  for c in ("species_code","occurrence","vantage","geog_site_area_cat3"))
        if any(not s for s in key):
            empty["MISSING_4_WAY_JOIN_KEY"]+=1
            continue
        if key in lookup:
            repeated["DUPLICATE_4_WAY_KEY"]+=1
            lookup[key]=None  # multiple original rules => ambiguous
        else:
            lookup[key]=row
    return lookup,{"certainty_missing_original_keys":dict(empty),
                   "certainty_duplicate_original_keys":dict(repeated)}


def rows_linked(sources):
    geog=index_geography(sources["potential_geography_2025"])
    old_ids={source_id(x.get("Geog_Site_ID"))
             for x in sources["potential_geography_2021"] if source_id(x.get("Geog_Site_ID"))}
    certify,cert_qa=build_certainty(sources["survey_certainty_rules"])
    all_surveys=[]
    missing_geo=Counter()
    lookup_results=Counter()
    statuses=Counter()
    species=Counter()
    vantage=Counter()
    expert=defaultdict(Counter)
    temporal=Counter()
    duplicates=defaultdict(list)
    for source_name in ("penguin_original_search","flying_seabird_original_search"):
        for ix,original in enumerate(sources[source_name]):
            site=source_id(original.get("geog_site_id"))
            species_name=source_text(original.get("species_code")).upper()
            occurrence=source_text(original.get("occurrence"))
            vantage_orig=source_text(original.get("vantage"))
            statuses[(source_name,occurrence or "<MISSING>")]+=1
            species[species_name or "<MISSING>"]+=1
            vantage[vantage_orig or "<MISSING>"]+=1
            for n in ("species_expert_validation","site_expert_validation",
                      "occurrence_expert_validation"):
                expert[n][source_text(original.get(n)) or "<MISSING>"]+=1
            geographic=geog.get(site)
            if geographic is None:
                missing_geo[source_name]+=1
            cat=source_text(geographic.get("geog_site_area_cat3")) if geographic else ""
            lookup=(species_name.casefold(),occurrence.casefold(),
                    vantage_orig.casefold(),cat.casefold())
            certainty=certify.get(lookup,"NO_CERTAINTY_SOURCE_MATCH")
            cflag=("EXACT_CERTAINTY_RULE" if isinstance(certainty,dict)
                   else "AMBIGUOUS_CERTAINTY_RULE" if certainty is None
                   else "NO_CERTAINTY_SOURCE_MATCH")
            lookup_results[cflag]+=1
            season=source_text(original.get("season"))
            year=source_text(original.get("year"))
            if season:temporal["WITH_SOURCE_SEASON"]+=1
            else:temporal["NO_SOURCE_SEASON"]+=1
            if year:temporal["WITH_SOURCE_YEAR"]+=1
            else:temporal["NO_SOURCE_YEAR"]+=1
            if site and season and species_name:
                duplicates[(site,species_name,season)].append((occurrence,source_name))
            entry={
                "source_study":source_name,
                "geog_site_id":site,
                "sub_group":source_text(original.get("sub_group")),
                "species_code":species_name,
                "occurrence":occurrence,
                "vantage":vantage_orig,
                "day":source_text(original.get("day")),
                "day_uncertainty":source_text(original.get("day_uncertainty")),
                "month":source_text(original.get("month")),
                "year":year,
                "season":season,
                "obs_number":source_text(original.get("obs_number")),
                "breed_site_id":source_text(original.get("breed_site_id")),
                "species_expert_validation":source_text(original.get("species_expert_validation")),
                "site_expert_validation":source_text(original.get("site_expert_validation")),
                "occurrence_expert_validation":source_text(original.get("occurrence_expert_validation")),
                "source_geog_site_lookup":"EXACT_2025_SITE_ID" if geographic else "NO_2025_SITE_MATCH",
                "source_geog_site_feature":source_text(geographic.get("geog_site_feature")) if geographic else "",
                "source_geog_site_area":source_text(geographic.get("geog_site_area")) if geographic else "",
                "source_geog_site_area_cat3":cat,
                "source_latitude":source_text(geographic.get("geog_site_latitude")) if geographic else "",
                "source_longitude":source_text(geographic.get("geog_site_longitude")) if geographic else "",
                "source_distance_station":source_text(geographic.get("distance_station")) if geographic else "",
                "source_distance_inland_coast":source_text(geographic.get("distance_inland_coast")) if geographic else "",
                "certainty_lookup":cflag,
                "rel_sp_detect":source_text(certainty.get("rel_sp_detect")) if isinstance(certainty,dict) else "",
                "rel_vantage_detect":source_text(certainty.get("rel_vantage_detect")) if isinstance(certainty,dict) else "",
                "site_search_completeness":source_text(certainty.get("site_search_completeness")) if isinstance(certainty,dict) else "",
                "certainty_product":source_text(certainty.get("product")) if isinstance(certainty,dict) else "",
                "occurrence_cert_cat":source_text(certainty.get("occurrence_cert_cat")) if isinstance(certainty,dict) else "",
                "scientific_absence_is_not_inferred_from_nonreporting":True,
            }
            all_surveys.append(entry)
    dup_pairs={key:reports for key,reports in duplicates.items() if len(reports)>1}
    contradictory=sum(len(set(st for st,source in reports))>1 for reports in dup_pairs.values())
    geo_features=Counter(source_text(g.get("geog_site_feature")) or "<EMPTY>" for g in geog.values())
    geo_area_cat=Counter(source_text(g.get("geog_site_area_cat3")) or "<EMPTY>" for g in geog.values())
    area_nonmissing=0
    area_nonpositive=0
    for row in geog.values():
        area=source_text(row.get("geog_site_area"))
        if not area:continue
        area_nonmissing+=1
        try:
            value=float(area)
        except ValueError:
            area_nonpositive+=1
            continue
        if not math.isfinite(value) or value<=0:area_nonpositive+=1
    geo_rows=[{
        "geog_site_id":site,
        "sub_group":source_text(g.get("sub_group")),
        "geog_site_feature":source_text(g.get("geog_site_feature")),
        "geog_site_area":source_text(g.get("geog_site_area")),
        "geog_site_area_cat3":source_text(g.get("geog_site_area_cat3")),
        "latitude":source_text(g.get("geog_site_latitude")),
        "longitude":source_text(g.get("geog_site_longitude")),
        "distance_station":source_text(g.get("distance_station")),
        "distance_inland_coast":source_text(g.get("distance_inland_coast")),
        "was_in_2021_original_geography":site in old_ids,
    } for site,g in sorted(geog.items())]
    result={
        "status":"V13_ORIGINAL_AADC_8_SEABIRD_SURVEY_LINK_AND_COVERAGE_NO_EFFECT",
        "original_2025_geographic_site_count":len(geo_rows),
        "original_2021_spatial_reference_site_count":len(old_ids),
        "new_2025_siteIDs_absent_from_2021":len(set(geog)-old_ids),
        "geog_siteIDs_2021_absent_from_2025":len(old_ids-set(geog)),
        "geog_siteIDs_shared_across_original_versions":len(set(geog)&old_ids),
        "original_2025_site_feature_classes":dict(geo_features),
        "original_2025_site_area_category_values":dict(geo_area_cat),
        "original_2025_site_area_present":area_nonmissing,
        "original_2025_nonpositive_nonfinite_or_unparsable_site_areas":area_nonpositive,
        "original_eight_seabird_source_observations":len(all_surveys),
        "source_search_occurrence_categories":{f"{k[0]}:{k[1]}":v for k,v in sorted(statuses.items())},
        "source_species_code_counts":dict(species),
        "source_vantage_categories":dict(vantage),
        "original_expert_validation_categories":{k:dict(v) for k,v in expert.items()},
        "source_certainty_lookup_state":dict(lookup_results),
        "original_certainty_key_QA":cert_qa,
        "source_breeding_year_fields":dict(temporal),
        "source_observations_unmatched_to_2025_geography":dict(missing_geo),
        "original_site_species_season_groups":len(duplicates),
        "source_site_species_season_groups_with_multiple_events":len(dup_pairs),
        "original_site_species_season_groups_with_conflicting_occurrence_categories":contradictory,
        "no_nonreporting_recoded_as_absence":True,
        "no_surveyed_zeros_fabricated":True,
        "ecological_effects_fitted":0,
    }
    return result,all_surveys,geo_rows


def emit_csv(path,rows,cols):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",newline="",encoding="utf-8") as output:
        writer=csv.DictWriter(output,fieldnames=cols,extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def main():
    p=argparse.ArgumentParser()
    p.add_argument("--out-json",type=Path,required=True)
    p.add_argument("--out-sites",type=Path,required=True)
    p.add_argument("--out-surveys",type=Path,required=True)
    args=p.parse_args()
    receipt={"status":"HOLD_AADC_SITE_OCCUPANCY_OR_CERTAINTY_SOURCE",
             "ecological_effects_fitted":0}
    try:
        with requests.Session() as session:
            manifest,index=public_dataset_manifests(session)
            sources,verified=fetch_original_sources(session,manifest)
        report,events,geography=rows_linked(sources)
        emit_csv(args.out_sites,geography,OUT_SITE_COLS)
        emit_csv(args.out_surveys,events,OUT_SURVEY_COLS)
        receipt={
            **report,
            "original_source_schemas_and_hashes":verified,
            "public_AADC_API_index_SHA256":index["source_sha256"],
            "derived_geographic_site_csv_SHA256":
                hashlib.sha256(args.out_sites.read_bytes()).hexdigest(),
            "derived_original_eight_seabird_search_csv_SHA256":
                hashlib.sha256(args.out_surveys.read_bytes()).hexdigest(),
            "original_source_DWCA_2023_APBP_not_mistaken_for_AADC_2025":True,
            "artifacts_store_public_derived_original_records_not_Git_commit":True,
        }
    except Exception as exc:
        receipt["error_type"]=type(exc).__name__
        receipt["error"]=str(exc)[:750]
    args.out_json.parent.mkdir(parents=True,exist_ok=True)
    args.out_json.write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
    if receipt["status"].startswith("HOLD"):
        raise SystemExit(2)


if __name__=="__main__":
    main()
