#!/usr/bin/env python3
"""Outcome-blind schema audit for East Antarctic penguin expansion sources.

Permitted: archive members, worksheet/table names, column headers, row counts,
file hashes, and identifier-column names. Prohibited: summaries of biological
values, occupancy-state frequencies, transitions, or ecological effects.
"""
from __future__ import annotations
import argparse, csv, hashlib, io, json, zipfile
from pathlib import Path
import openpyxl
import xlrd

def sha256(path: Path) -> str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda:f.read(1024*1024),b""): h.update(block)
    return h.hexdigest()

def csv_schema(raw: bytes, name: str) -> dict:
    text=None; encoding=None
    for enc in ("utf-8-sig","utf-8","latin-1"):
        try: text=raw.decode(enc); encoding=enc; break
        except UnicodeDecodeError: pass
    if text is None: return {"name":name,"kind":"csv","readable":False}
    rows=csv.reader(io.StringIO(text), delimiter="\t" if name.lower().endswith(".tsv") else ",")
    try: header=next(rows); n=sum(1 for _ in rows)
    except StopIteration: header=[]; n=0
    return {"name":name,"kind":"csv","encoding":encoding,
            "columns":[str(x).strip() for x in header],"data_rows":int(n)}

def xlsx_schema(raw: bytes, name: str) -> dict:
    wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True)
        try: first=next(it)
        except StopIteration: first=()
        sheets.append({"sheet":ws.title,
                       "columns":[None if v is None else str(v).strip() for v in first],
                       "max_row":int(ws.max_row or 0),"max_column":int(ws.max_column or 0)})
    return {"name":name,"kind":"xlsx","sheets":sheets}

def xls_schema(raw: bytes, name: str) -> dict:
    try:
        book=xlrd.open_workbook(file_contents=raw,on_demand=True)
    except Exception as exc:
        return {"name":name,"kind":"xls","readable":False,"reason":str(exc)}
    sheets=[]
    for sname in book.sheet_names():
        ws=book.sheet_by_name(sname)
        header=ws.row_values(0) if ws.nrows else []
        sheets.append({"sheet":sname,
                       "columns":[None if v in ("",None) else str(v).strip() for v in header],
                       "max_row":int(ws.nrows),"max_column":int(ws.ncols)})
    return {"name":name,"kind":"xls","sheets":sheets}

def inspect_bytes(raw: bytes, name: str) -> dict:
    low=name.lower()
    if low.endswith((".csv",".tsv")): return csv_schema(raw,name)
    if low.endswith(".xlsx"): return xlsx_schema(raw,name)
    if low.endswith(".xls"): return xls_schema(raw,name)
    if zipfile.is_zipfile(io.BytesIO(raw)):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            names=z.namelist()
            if "[Content_Types].xml" in names and any(n.startswith("xl/") for n in names):
                return xlsx_schema(raw,name)
            tables=[]
            for member in names:
                ml=member.lower()
                if ml.endswith((".csv",".tsv")): tables.append(csv_schema(z.read(member),member))
                elif ml.endswith(".xlsx"): tables.append(xlsx_schema(z.read(member),member))
                elif ml.endswith(".xls"): tables.append(xls_schema(z.read(member),member))
            return {"name":name,"kind":"zip","members":names,"tables":tables}
    return {"name":name,"kind":"opaque","readable_schema":False}

def inspect_file(path: Path) -> dict:
    raw=path.read_bytes()
    out={"path":str(path),"bytes":len(raw),"sha256":sha256(path),"prefix_hex":raw[:16].hex()}
    out.update(inspect_bytes(raw,path.name))
    return out

def inspect_dir(path: Path) -> dict:
    files=sorted(p for p in path.rglob("*") if p.is_file())
    return {"path":str(path),"n_files":len(files),"files":[inspect_file(p) for p in files]}

def main()->int:
    p=argparse.ArgumentParser()
    p.add_argument("--occupancy-2025-dir",required=True,type=Path)
    p.add_argument("--occupancy-2016-dir",required=True,type=Path)
    p.add_argument("--plos-s1",required=True,type=Path)
    p.add_argument("--out",required=True,type=Path)
    a=p.parse_args()
    result={"schema_version":2,"analysis_id":"east-antarctica-expansion-schema-audit-v2",
            "status":"outcome_blind_schema_only",
            "sources":{"occupancy_2025":inspect_dir(a.occupancy_2025_dir),
                       "occupancy_2016":inspect_dir(a.occupancy_2016_dir),
                       "plos_s1":inspect_file(a.plos_s1)},
            "biological_values_summarized":False,
            "occupancy_state_frequencies_computed":False,
            "colonization_events_counted":False}
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.out.write_text(json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps(result,indent=2,sort_keys=True))
    return 0
if __name__=="__main__": raise SystemExit(main())
