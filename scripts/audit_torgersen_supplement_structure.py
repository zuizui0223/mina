#!/usr/bin/env python3
"""Outcome-blind structural audit of the Torgersen supplementary DOCX.

Reads DOCX package structure, table dimensions/header rows, table captions and
embedded-object names. It does not export body-table rows or calculate any
ecological result.
"""
from __future__ import annotations
import argparse, json, re, zipfile
from pathlib import Path
from docx import Document

def audit_docx(path: Path) -> dict:
    doc=Document(path)
    tables=[]
    for i,t in enumerate(doc.tables,1):
        header=[]
        if t.rows:
            header=[c.text.strip() for c in t.rows[0].cells]
        tables.append({
            "table_index":i,
            "rows":len(t.rows),
            "columns":len(t.columns),
            "header":header,
        })
    captions=[]
    for p in doc.paragraphs:
        x=" ".join(p.text.split())
        if not x: continue
        if re.match(r"^(Table|Fig(?:ure)?|Supplement|S\d+)",x,re.I):
            captions.append(x[:500])
    with zipfile.ZipFile(path) as z:
        names=z.namelist()
    return {
        "bytes":path.stat().st_size,
        "table_count":len(tables),
        "tables":tables,
        "captions":captions,
        "embedded_objects":[n for n in names if "/embeddings/" in n],
        "media_count":sum("/media/" in n for n in names),
        "package_members":len(names),
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--supplement",required=True,type=Path)
    p.add_argument("--site-guide",type=Path)
    p.add_argument("--out",required=True,type=Path)
    a=p.parse_args()
    result={
      "schema_version":1,
      "analysis_id":"torgersen-supplement-structure-gate-v1",
      "status":"outcome_blind_structure_only",
      "supplement":audit_docx(a.supplement),
      "body_table_rows_exported":False,
      "ecological_effect_computed":False
    }
    if a.site_guide and a.site_guide.exists():
        txt=a.site_guide.read_text(encoding="utf-8",errors="replace")
        result["duke_site_guide"]={
          "bytes":a.site_guide.stat().st_size,
          "lines":[line for line in txt.splitlines() if line.strip()]
        }
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.out.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(result,indent=2,ensure_ascii=False))
if __name__=="__main__": main()
