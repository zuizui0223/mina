#!/usr/bin/env python3
"""Build Ecology Appendix S1 as a single DOCX/PDF-ready file."""
from __future__ import annotations

import argparse
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from mina.ecosphere_submission_metadata import (
    author_line,
    load_metadata,
    require_complete_metadata,
)


AUTHOR_PLACEHOLDER = "[SAME AUTHOR LIST AS MAIN MANUSCRIPT]"


def _page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _prepare_markdown(
    source: str,
    *,
    metadata: dict[str, object] | None,
) -> str:
    authors = AUTHOR_PLACEHOLDER if metadata is None else author_line(metadata)
    return source.replace(AUTHOR_PLACEHOLDER, authors)


def _set_table_widths(table, widths: list[float]) -> None:
    """Apply conservative fixed widths for portrait-page appendix tables."""
    table.autofit = False
    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def _format_docx(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.footer.is_linked_to_previous = False
        _page_number(section.footer.paragraphs[0])

    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for paragraph in doc.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:eastAsia"), "Times New Roman"
            )

    # Dense tables may use 10 pt and single spacing while retaining journal readability.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)
                        run._element.get_or_add_rPr().rFonts.set(
                            qn("w:eastAsia"), "Times New Roman"
                        )


def build(
    source_path: Path,
    caption_path: Path,
    figure_path: Path,
    out_docx: Path,
    *,
    metadata_path: Path | None = None,
    require_complete: bool = False,
) -> None:
    source = source_path.read_text(encoding="utf-8")
    caption_lines = [
        line for line in caption_path.read_text(encoding="utf-8").splitlines()
        if not line.startswith("# Supplementary figure caption")
    ]
    while caption_lines and not caption_lines[0].strip():
        caption_lines.pop(0)
    if not caption_lines or not caption_lines[0].startswith("## Figure S1."):
        raise ValueError("supplementary caption must begin with '## Figure S1.'")
    caption_heading = caption_lines.pop(0).removeprefix("## ").strip()
    caption_body = " ".join(line.strip() for line in caption_lines if line.strip())

    metadata = None
    if metadata_path is not None:
        metadata = load_metadata(metadata_path)
        if require_complete:
            require_complete_metadata(metadata)

    md_text = _prepare_markdown(source, metadata=metadata)

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        md = tmp / "appendix.md"
        raw = tmp / "appendix_raw.docx"
        md.write_text(md_text, encoding="utf-8")

        subprocess.run(
            [
                "pandoc",
                str(md),
                "--from=markdown+pipe_tables+tex_math_single_backslash+tex_math_dollars",
                "--to=docx",
                "-o",
                str(raw),
            ],
            check=True,
        )

        doc = Document(raw)
        _format_docx(doc)

        # Keep the wide regional inference table together on a fresh page.
        for paragraph in doc.paragraphs:
            if paragraph.text.strip().startswith("Table S4."):
                paragraph.paragraph_format.page_break_before = True
                paragraph.paragraph_format.keep_with_next = True
                break

        # Five appendix tables are emitted in order S1-S5.
        if len(doc.tables) >= 5:
            _set_table_widths(
                doc.tables[3],
                [1.10, 0.55, 0.75, 0.45, 0.55, 0.55, 0.80, 0.45, 0.70],
            )
            _set_table_widths(
                doc.tables[4],
                [1.00, 0.65, 0.95, 0.95, 0.90, 0.90],
            )

        # Figure S1 is part of Appendix S1 and must not be uploaded separately.
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(figure_path), width=Inches(6.5))
        cap = doc.add_paragraph()
        cap.paragraph_format.line_spacing = 2.0
        lead = cap.add_run(caption_heading + " ")
        lead.bold = True
        cap.add_run(caption_body)

        props = doc.core_properties
        props.title = "Appendix S1 - Breeding-space contraction recurs across spatial scales in Antarctic penguins"
        props.subject = "Ecology Supporting Information"
        props.author = ""
        props.last_modified_by = ""

        out_docx.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_docx)


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--source", required=True, type=Path)
    p.add_argument("--caption", required=True, type=Path)
    p.add_argument("--figure", required=True, type=Path)
    p.add_argument("--out-docx", required=True, type=Path)
    p.add_argument("--metadata", type=Path)
    p.add_argument("--require-complete-metadata", action="store_true")
    a = p.parse_args()

    build(
        a.source,
        a.caption,
        a.figure,
        a.out_docx,
        metadata_path=a.metadata,
        require_complete=a.require_complete_metadata,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
