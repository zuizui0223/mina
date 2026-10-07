#!/usr/bin/env python3
"""Build Supporting Information DOCX for the spatial-recovery Ecology submission.

Production-only builder. Scientific text and figure images are read from frozen
or freeze-compliant source files; this script does not alter scientific content.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FIGURES = [
    ("Figure S1", "figure3_ross_path_state_calibration.png"),
    ("Figure S2", "figure4_process_contrasts.png"),
]


def _set_layout(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def _page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _format_paragraph(paragraph, *, double: bool = True) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 2.0 if double else 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _extract_sections(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    heading: str | None = None
    body: list[str] = []
    for line in text.splitlines():
        if line.startswith("# ") and not line.startswith("## "):
            continue
        if line.startswith("## "):
            if heading is not None:
                sections.append((heading, "\n".join(body).strip()))
            heading = line[3:].strip()
            body = []
        else:
            body.append(line)
    if heading is not None:
        sections.append((heading, "\n".join(body).strip()))
    return sections


def _clean_inline(text: str) -> str:
    return text.replace("**", "")


def build(supplement_path: Path, figure_dir: Path, out_path: Path) -> None:
    source = supplement_path.read_text(encoding="utf-8")
    sections = _extract_sections(source)

    doc = Document()
    for section in doc.sections:
        _set_layout(section)
        section.footer.is_linked_to_previous = False
        _page_number(section.footer.paragraphs[0])

    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Supporting Information")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Population recovery can retrace local losses while spatial structure remains altered")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    section_map = {h: b for h, b in sections}

    for index, (label, filename) in enumerate(FIGURES):
        heading = next((h for h in section_map if h.startswith(label + ".")), None)
        if heading is None:
            raise ValueError(f"missing supplement caption for {label}")
        caption = section_map[heading]

        if index:
            doc.add_page_break()

        p = doc.add_paragraph()
        r = p.add_run(heading)
        r.bold = True
        _format_paragraph(p)

        for block in [b for b in caption.split("\n\n") if b.strip()]:
            if block.startswith("**Source image:**"):
                continue
            p = doc.add_paragraph(_clean_inline(block.replace("\n", " ")))
            _format_paragraph(p)

        image = figure_dir / filename
        if not image.exists():
            raise FileNotFoundError(image)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image), width=Inches(6.7))

    boundary = section_map.get("Inferential boundary")
    if boundary:
        doc.add_page_break()
        p = doc.add_paragraph()
        r = p.add_run("Inferential boundary")
        r.bold = True
        _format_paragraph(p)
        for block in [b for b in boundary.split("\n\n") if b.strip()]:
            p = doc.add_paragraph(_clean_inline(block.replace("\n", " ")))
            _format_paragraph(p)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--supplement", required=True, type=Path)
    p.add_argument("--figure-dir", required=True, type=Path)
    p.add_argument("--out", required=True, type=Path)
    a = p.parse_args()
    build(a.supplement, a.figure_dir, a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
