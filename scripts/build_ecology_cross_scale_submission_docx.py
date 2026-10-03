#!/usr/bin/env python3
"""Build an Ecology Article cross-scale Word Main Document preview/submission file."""
from __future__ import annotations

import argparse
import copy
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from mina.ecosphere_submission_metadata import (
    affiliation_lines,
    ai_disclosure_suffix,
    author_line,
    corresponding_text,
    load_metadata,
    present_address_text,
    require_complete_metadata,
)

MARKER = "[[SECTION_BREAK_AFTER_TITLE]]"

KEYWORDS = (
    "abundance–occupancy; Antarctica; colonial breeding; effective number; "
    "penguins; population decline; spatial contraction; spatial scale"
)

OPEN_RESEARCH = (
    "The Palmer Station Antarctica Long Term Ecological Research Adélie "
    "penguin census is publicly available at DOI "
    "10.6073/pasta/89dd52217ca37e3a72a67f7a9bc3c82e. The Signy Island "
    "Adélie penguin data are publicly available at DOI "
    "10.5285/daf2c4fd-c1e3-4e65-851f-d11f02c5b69d, and the Signy Island "
    "chinstrap penguin data at DOI "
    "10.5285/d0633a9b-8c56-4ae5-88bf-6ec6ba9017b9. The MAPPPD regional "
    "extension uses the public CCheCastaldo/mapppdr repository pinned at "
    "commit 88c73a507e0921b2541c218c71eaf16721bc6502. Analysis code, frozen "
    "contracts, provenance records, result receipts and reproducible figure "
    "workflows are publicly accessible for peer review at "
    "https://github.com/zuizui0223/mina. If accepted, the exact analysis/code "
    "release and derived outputs will be deposited in a permanent repository "
    "and the final Open Research Statement updated with that DOI."
)

AI_ACK = (
    "OpenAI ChatGPT (GPT-5.6 Sol) was used during analysis and manuscript "
    "development to assist with code drafting and review, literature searching, "
    "statistical sensitivity-analysis scripting and editorial drafting. All "
    "analyses were executed from version-controlled code, numerical outputs "
    "were checked against frozen result receipts and public source-data "
    "checksums, cited literature was independently verified, and the authors "
    "remain responsible for all analyses, interpretations and text."
)


def _extract_title(text: str) -> str:
    first = text.splitlines()[0] if text.splitlines() else ""
    if not first.startswith("# "):
        raise ValueError("manuscript must begin with an H1 title")
    title = first[2:].strip()
    if not title:
        raise ValueError("empty manuscript title")
    return title


def _strip_manuscript_source(text: str) -> str:
    lines = text.splitlines()
    out: list[str] = []
    skipping_refs = False
    for i, line in enumerate(lines):
        if i == 0 and line.startswith("# "):
            continue
        if line.startswith("**Ecology Article candidate") or line.startswith("**Ecology Report candidate"):
            continue
        if line.startswith("**Keywords:**"):
            continue
        if line.strip() == "## References":
            skipping_refs = True
            continue
        if skipping_refs:
            continue
        out.append(line)
    return "\n".join(out).rstrip()


def _strip_caption_heading(text: str) -> str:
    lines = text.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and lines[0].startswith("# ") and not lines[0].startswith("## "):
        lines.pop(0)
    while lines and not lines[0].strip():
        lines.pop(0)
    return "\n".join(lines).strip()


def _open_research(metadata: dict[str, object] | None) -> str:
    if metadata is None:
        return OPEN_RESEARCH
    review_url = str(metadata.get("review_code_url") or "").strip()
    doi = str(metadata.get("permanent_archive_doi") or "").strip()
    text = OPEN_RESEARCH.replace(
        "https://github.com/zuizui0223/mina",
        review_url or "https://github.com/zuizui0223/mina",
    )
    if doi:
        text = text.replace(
            "If accepted, the exact analysis/code release and derived outputs "
            "will be deposited in a permanent repository and the final Open "
            "Research Statement updated with that DOI.",
            "The exact analysis/code release and derived outputs are archived "
            "at DOI " + doi + ".",
        )
    return text


def _metadata_blocks(metadata: dict[str, object] | None) -> dict[str, str]:
    if metadata is None:
        return {
            "authors": "[AUTHOR 1], [AUTHOR 2], [...]",
            "affiliations": "[AFFILIATIONS — REQUIRED]",
            "present": "[PRESENT ADDRESSES OR DELETE]",
            "corresponding": "[ONE AUTHOR NAME], [EMAIL ADDRESS]",
            "funding": "[FUNDING ACKNOWLEDGMENTS AND GRANT IDENTIFIERS — REQUIRED]",
            "additional_ack": "",
            "contributions": "[AUTHOR CONTRIBUTIONS — REQUIRED; confirm CRediT roles with every author.]",
            "coi": "[CONFLICT-OF-INTEREST STATEMENT — REQUIRED; confirm with every author.]",
            "ai_suffix": "",
        }

    return {
        "authors": author_line(metadata),
        "affiliations": affiliation_lines(metadata),
        "present": present_address_text(metadata),
        "corresponding": corresponding_text(metadata),
        "funding": str(metadata.get("funding_acknowledgments") or "").strip(),
        "additional_ack": str(metadata.get("additional_acknowledgments") or "").strip(),
        "contributions": str(metadata.get("author_contributions") or "").strip(),
        "coi": str(metadata.get("conflict_of_interest") or "").strip(),
        "ai_suffix": ai_disclosure_suffix(metadata),
    }


def _sentence(text: str) -> str:
    value = text.strip()
    if not value:
        return ""
    return value if value[-1] in ".!?" else value + "."


def _combined_markdown(
    manuscript: str,
    captions: str,
    metadata: dict[str, object] | None,
) -> str:
    title = _extract_title(manuscript)
    body = _strip_manuscript_source(manuscript)
    caption_body = _strip_caption_heading(captions)
    blocks = _metadata_blocks(metadata)
    open_research = _open_research(metadata)

    present_line = (
        f"**Present address(es), if applicable:** {blocks['present']}\n\n"
        if metadata is None or blocks["present"] != "None."
        else ""
    )

    title_page = f"""Ecology

Article

# {title}

**Authors:** {blocks["authors"]}

**Affiliations:**  
{blocks["affiliations"]}

{present_line}**Corresponding author:** {blocks["corresponding"]}

## Open Research Statement

{open_research}

## Key words/phrases

{KEYWORDS}

{MARKER}

"""

    acknowledgments = " ".join(
        x for x in (
            _sentence(blocks["funding"]),
            _sentence(blocks["additional_ack"]),
            _sentence(AI_ACK + blocks["ai_suffix"]),
        )
        if x
    )

    backmatter = f"""

## Acknowledgments

{acknowledgments}

## Author Contributions

{blocks["contributions"]}

## Conflict of Interest Statement

{blocks["coi"]}

## References

::: {{#refs}}
:::

## Figure captions

{caption_body}
"""
    return title_page + body + backmatter


def _set_layout(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def _remove_line_numbering(sect_pr) -> None:
    for node in list(sect_pr.findall(qn("w:lnNumType"))):
        sect_pr.remove(node)


def _add_line_numbering(
    sect_pr,
    *,
    start: str = "0",
    restart: str = "newSection",
) -> None:
    _remove_line_numbering(sect_pr)
    node = OxmlElement("w:lnNumType")
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:start"), start)
    node.set(qn("w:restart"), restart)
    node.set(qn("w:distance"), "360")
    sect_pr.append(node)


def _suppress_line_number(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:suppressLineNumbers")) is None:
        p_pr.append(OxmlElement("w:suppressLineNumbers"))


def _page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _suppress_line_number(paragraph)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _split_title_section(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    marker_index = next(
        (i for i, p in enumerate(paragraphs) if p.text.strip() == MARKER),
        None,
    )
    if marker_index is None:
        raise ValueError("title/body section marker missing")

    marker = paragraphs[marker_index]
    boundary = next(
        (p for p in reversed(paragraphs[:marker_index]) if p.text.strip()),
        None,
    )
    if boundary is None:
        raise ValueError("title-page boundary paragraph missing")

    body_sect_pr = doc.element.body.sectPr
    title_sect_pr = copy.deepcopy(body_sect_pr)

    # Ecology requires continuous line numbering after the title page.
    _remove_line_numbering(title_sect_pr)
    _add_line_numbering(body_sect_pr, start="0", restart="newSection")

    type_node = title_sect_pr.find(qn("w:type"))
    if type_node is None:
        type_node = OxmlElement("w:type")
        title_sect_pr.insert(0, type_node)
    type_node.set(qn("w:val"), "nextPage")

    boundary_p_pr = boundary._p.get_or_add_pPr()
    existing = boundary_p_pr.find(qn("w:sectPr"))
    if existing is not None:
        boundary_p_pr.remove(existing)
    boundary_p_pr.append(title_sect_pr)

    parent = marker._p.getparent()
    if parent is None:
        raise ValueError("section marker is detached")
    parent.remove(marker._p)


def _apply_format(doc: Document, title: str) -> None:
    for section in doc.sections:
        _set_layout(section)

    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def format_paragraph(paragraph, spacing: float) -> None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = spacing
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:eastAsia"), "Times New Roman"
            )

    body_started = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Abstract":
            body_started = True
        if not body_started:
            _suppress_line_number(paragraph)
        format_paragraph(paragraph, 2.0 if body_started else 1.0)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph, 2.0)

    for i, section in enumerate(doc.sections):
        section.footer.is_linked_to_previous = False
        _page_number(section.footer.paragraphs[0])
        if i == 0:
            _remove_line_numbering(section._sectPr)
        else:
            _add_line_numbering(
                section._sectPr,
                start="1",
                restart="newSection",
            )

    props = doc.core_properties
    props.title = title
    props.subject = "Ecology Article submission"
    props.author = ""
    props.last_modified_by = ""


def build(
    manuscript_path: Path,
    bibliography_path: Path,
    captions_path: Path,
    out_path: Path,
    metadata_path: Path | None = None,
    require_complete: bool = False,
) -> None:
    manuscript = manuscript_path.read_text(encoding="utf-8")
    captions = captions_path.read_text(encoding="utf-8")
    metadata = None
    if metadata_path is not None:
        metadata = load_metadata(metadata_path)
        if require_complete:
            require_complete_metadata(metadata)

    title = _extract_title(manuscript)
    source = _combined_markdown(manuscript, captions, metadata)

    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        md = tmp / "submission.md"
        raw = tmp / "raw.docx"
        split = tmp / "split.docx"
        md.write_text(source, encoding="utf-8")

        subprocess.run(
            [
                "pandoc",
                str(md),
                "--from=markdown+raw_attribute+tex_math_single_backslash+tex_math_dollars",
                "--to=docx",
                "--citeproc",
                f"--bibliography={bibliography_path}",
                "-o",
                str(raw),
            ],
            check=True,
        )

        doc = Document(raw)
        _split_title_section(doc)
        doc.save(split)

        doc = Document(split)
        _apply_format(doc, title)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--manuscript", required=True, type=Path)
    p.add_argument("--bibliography", required=True, type=Path)
    p.add_argument("--captions", required=True, type=Path)
    p.add_argument("--out", required=True, type=Path)
    p.add_argument("--metadata", type=Path)
    p.add_argument("--require-complete-metadata", action="store_true")
    a = p.parse_args()
    build(
        a.manuscript,
        a.bibliography,
        a.captions,
        a.out,
        metadata_path=a.metadata,
        require_complete=a.require_complete_metadata,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
