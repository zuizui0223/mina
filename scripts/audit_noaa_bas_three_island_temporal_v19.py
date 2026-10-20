#!/usr/bin/env python3
"""V19 pinned, source-only three-island Gentoo year-matched data audit. No count magnitudes."""
import argparse
from collections import Counter,defaultdict
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re
from openpyxl import load_workbook
import audit_noaa_bas_independent_island_sources_v18 as v18

NOAA_SHA="a57370f82adf82b53da11ce5d85dada1ae3f3750ef246312b59d0f168763eef4"
BAS_SHA="53fe3821f313ed35840376fcf4d11c8e0e8fbbd63f01ba84a5e21c2db4df6d6b"
NOAA_COLS=("Copa_Adelie","Copa_Gentoo","Copa_Chinstrap",
           "CapeShirreff_Chinstrap","CapeShirreff_Gentoo")
BAS_COL="TOTAL_NUMBER_OF_PAIRS"

def source_year(raw):
    s=str(raw).strip()
    if re.fullmatch(r"(?:19|20)\d{2}",s):
        return int(s)
    m=re.fullmatch(r"((?:19|20)\d{2})\s*[-/]\s*(\d{2}|\d{4})",s)
    if m:
        first=int(m.group(1))
        end=int(m.group(2))
        if end==(first+1)%100 or end==first+1:
            return first
    return None

def lexical_numeric(v):
    if isinstance(v,bool) or v is None:
        return False
    if isinstance(v,(int,float)):
        return math.isfinite(v)
    text=str(v).strip()
    # lexical recognition, no use of value or count magnitude.
    return bool(re.fullmatch(r"[+-]?(?:\d+(?:\.\d+)?|\.\d+)",text))

def extract_noaa(raw):
    w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    try:
        if w.sheetnames!=["census data"]:
            raise ValueError("SOURCE_STOP NOAA exact original worksheet changed")
        ws=w["census data"]
        it=ws.iter_rows(values_only=True)
        head=list(next(it))
        if [str(v).strip() for v in head[:6]]!=["YEAR",*NOAA_COLS]:
            raise ValueError("SCHEMA_STOP NOAA exact species columns not stable")
        years=defaultdict(list)
        unsupported=[]
        for row in it:
            if not row or all(v is None or str(v).strip()=="" for v in row):
                continue
            raw_year=row[0] if row else None
            year=source_year(raw_year)
            if year is None or not 1977<=year<=2015:
                unsupported.append(str(raw_year)[:60])
                continue
            years[year].append({col:lexical_numeric(row[k+1]) if k+1<len(row) else False
                                for k,col in enumerate(NOAA_COLS)})
        valid={year:records[0] for year,records in years.items() if len(records)==1}
        duplicates={str(year):len(records) for year,records in years.items() if len(records)>1}
        return valid,{
            "original_publisher_season_rows":sum(map(len,years.values()))+len(unsupported),
            "parsed_distinct_season_start_years":len(years),
            "duplicate_year_groups_excluded":duplicates,
            "unsupported_year_labels":unsupported[:10],
            "first_last_parsed_year":([min(years),max(years)] if years else None),
            "numeric_lexical_valid_years_by_species":{
                col:sum(one[col] for one in valid.values()) for col in NOAA_COLS},
            "missing_nonnumeric_years_by_species":{
                col:len(valid)-sum(one[col] for one in valid.values()) for col in NOAA_COLS},
        }
    finally:
        w.close()

def extract_bas(raw):
    try:
        txt=raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        txt=raw.decode("latin-1")
    data=list(csv.DictReader(io.StringIO(txt)))
    if not data or BAS_COL not in data[0] or "SEASON" not in data[0]:
        raise ValueError("SCHEMA_STOP BAS original gentoo season/pair columns changed")
    years=defaultdict(list)
    unsupported=[]
    for row in data:
        start=source_year(row.get("SEASON"))
        if start is None or not 1978<=start<=2020:
            unsupported.append(str(row.get("SEASON"))[:60])
            continue
        years[start].append(lexical_numeric(row.get(BAS_COL)))
    valid={year:vals[0] for year,vals in years.items() if len(vals)==1}
    return valid,{
        "original_rows":len(data),
        "distinct_parsed_original_seasons":len(years),
        "duplicate_year_groups_excluded":{str(y):len(vals) for y,vals in years.items() if len(vals)>1},
        "unsupported_year_labels":unsupported[:12],
        "first_last_parsed_year":([min(years),max(years)] if years else None),
        "original_BAS_TOTAL_NUMBER_OF_PAIRS_numeric_lexical_years":sum(valid.values()),
        "original_BAS_TOTAL_NUMBER_OF_PAIRS_nonnumeric_or_missing_years":len(valid)-sum(valid.values()),
        "original_source_column_is_not_verified_NOAA_count_equivalent":True,
    }

def compare_sources(noaa,bas):
    sets={
        "KingGeorge_Copa_Gentoo":{yr for yr,vals in noaa.items() if vals["Copa_Gentoo"]},
        "Livingston_CapeShirreff_Gentoo":{yr for yr,vals in noaa.items() if vals["CapeShirreff_Gentoo"]},
        "Signy_BAS_Gentoo_Pairs":{yr for yr,good in bas.items() if good},
    }
    years=set.intersection(*sets.values())
    runs=sorted(y for y in years if y+1 in years)
    shared_noaa={
        "Copacabana_ADPE_CHPE_GEPE":sorted(y for y,r in noaa.items() if
           r["Copa_Adelie"] and r["Copa_Gentoo"] and r["Copa_Chinstrap"]),
        "CapeShirreff_CHPE_GEPE":sorted(y for y,r in noaa.items() if
           r["CapeShirreff_Gentoo"] and r["CapeShirreff_Chinstrap"]),
    }
    source_gate=(len(years)>=10 and len(runs)>=8)
    return {
        "all_three_island_gentoo_numeric_lexical_common_seasons":len(years),
        "all_three_gentoo_common_start_years":sorted(years),
        "all_three_island_gentoo_consecutive_year_windows":len(runs),
        "all_three_gentoo_consecutive_start_years":runs,
        "noaa_aligned_multispecies_seasons_by_local_monitoring_site":{
            name:len(ys) for name,ys in shared_noaa.items()},
        "prespecified_three_island_temporal_source_support_pass":source_gate,
        "source_count_units_fully_harmonized":False,
        "population_density_or_growth_effect_fitted":False,
        "physical_island_n":3,
        "site_systems_are_not_whole_island_demographic_censuses":True,
        "scientific_decision":("HOLD_NOAA_BAS_CENSUS_STAGE_DEFINITION_AND_INDEPENDENT_ENVIRONMENT"
             if source_gate else "STOP_INSUFFICIENT_THREE_ISLAND_YEAR_OVERLAP"),
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--out",type=Path,required=True)
    args=p.parse_args()
    report={"status":"HOLD_ORIGINAL_SOURCE","ecological_effect_fitted":False}
    try:
        noaam,receipt_noaa=v18.source_get(v18.NOAA_URL,"noaa")
        basm,receipt_bas=v18.source_get(v18.BAS_URL,"bas")
        if receipt_noaa["sha256"]!=NOAA_SHA or receipt_bas["sha256"]!=BAS_SHA:
            raise ValueError("SOURCE_STOP NOAA/BAS pinned original SHA256 drift")
        n,nmeta=extract_noaa(noaam)
        b,bmeta=extract_bas(basm)
        report={
            "status":"V19_ORIGINAL_THREE_ISLAND_SEASON_SUPPORT_ONLY",
            "noaa_original_sha256":NOAA_SHA,"bas_original_sha256":BAS_SHA,
            "noaa_structural_coverage":nmeta,"bas_gentoo_structural_coverage":bmeta,
            "comparison":compare_sources(n,b),
            "biological_count_magnitudes_opened":0,
            "biology_regression_fit":0,
        }
    except Exception as exc:
        report["error_type"]=type(exc).__name__
        report["error"]=str(exc)[:700]
    args.out.parent.mkdir(parents=True,exist_ok=True)
    args.out.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))
    if report["status"]=="HOLD_ORIGINAL_SOURCE":
        raise SystemExit(2)
if __name__=="__main__":
    main()
