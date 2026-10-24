#!/usr/bin/env python3
"""V32: exact AADC original regional bootstrap workbook HEADER/grain audit.

Never reads original animal-count, population estimate or bootstrap-draw values
for analysis or in machine-readable receipts.
"""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from urllib.parse import quote
import requests
from openpyxl import load_workbook
import audit_aadc_three_island_original_v28 as v28

UUID="7673853e-9360-4563-a0e9-324300cfa0b7"
OBJECT="W Mac. Robertson Land Adelie penguin population estimates.xlsx"
EXPECTED_BYTES=51218
EXPECTED_SOURCE_DOI="10.26179/s2qa-s344"
GATE_MIN_PHYSICAL_ISLANDS=10
GATE_MIN_REGIONS=3
IDENTIFIER_FIELDS=(
    "island","island name","island id","island_name","island_id",
    "site","site name","site id","site_name","site_id",
    "colony","colony name","colony id","colony_name","colony_id",
    "region","region name","region id","region_name","region_id",
    "island group","island_group"
)
HINTS=("island","site","colony","region","bootstrap","estimated","estimate",
       "population","nested","nest","occupied","2009","2019","percentile","median")
LIMIT_ROWS_PER_SHEET=8
MAX_SHEETS=30
MAX_HEADER_COLUMNS=45

def safe_label(value):
    if not isinstance(value,str):
        return None
    s=value.strip()
    if not s:
        return None
    if re.fullmatch(r"[-+]?(?:\d+(?:\.\d+)?|\.\d+)(?:[Ee][-+]?\d+)?",s):
        return "[NUMERIC_TEXT_REDACTED]"
    # Text label only, not original number fields or raw count magnitudes.
    return s[:115]

def metadata_sheet(sheet):
    rows=[]
    labels=[]
    for index,row in enumerate(sheet.iter_rows(min_row=1,max_row=LIMIT_ROWS_PER_SHEET,
                                               values_only=True),start=1):
        text=[]
        for val in row[:MAX_HEADER_COLUMNS]:
            label=safe_label(val)
            if label:
                text.append(label)
                if label!="[NUMERIC_TEXT_REDACTED]":
                    labels.append(label)
        if text:
            rows.append({"row_number":index,"source_TEXT_labels_only":text})
    ids=[]
    for label in labels:
        clean=" ".join(re.split(r"\s+",label.strip().casefold()))
        if clean in IDENTIFIER_FIELDS:
            ids.append(label)
    hints=[text for text in labels if any(q in text.casefold() for q in HINTS)]
    return {
        "source_sheet_name":sheet.title,
        "declared_source_rows":int(sheet.max_row or 0),
        "declared_source_columns":int(sheet.max_column or 0),
        "literal_text_in_first_eight_rows":rows,
        "exact_geographic_identifier_text_labels":sorted(set(ids)),
        "source_text_hints":list(dict.fromkeys(hints))[:25],
        "source_population_count_cells_read":0,
        "bootstrap_draws_opened":0,
        "original_row_counts_not_independent_islands":True,
    }

def inspect(raw):
    if not raw.startswith(b"PK"):
        raise ValueError("SOURCE_STOP AADC publisher member is not Excel OOXML")
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    try:
        if not wb.sheetnames or len(wb.sheetnames)>MAX_SHEETS:
            raise ValueError("SCHEMA_STOP invalid original workbook sheet list")
        sheets=[metadata_sheet(sheet) for sheet in wb.worksheets]
        candidate=sorted(set(q for row in sheets for q in
                             row["exact_geographic_identifier_text_labels"]))
        hints=[v.casefold() for row in sheets for v in row["source_text_hints"]]
        return {
            "original_publisher_workbook_sheets":sheets,
            "literal_original_geography_field_names":candidate,
            "publisher_text_mentions_2009":any("2009" in h for h in hints),
            "publisher_text_mentions_2019":any("2019" in h for h in hints),
            "original_sheets_with_island_name_id_header":sum(any(
                "island" in x.casefold() for x in row["exact_geographic_identifier_text_labels"])
                for row in sheets),
            "original_island_identity_rows_verified":0,
            "original_island_geography_crosswalk_verified":False,
            "original_geographically_distinct_physical_islands_verified":0,
            "frozen_next_gate_minimum_islands":GATE_MIN_PHYSICAL_ISLANDS,
            "frozen_next_gate_minimum_independent_regions":GATE_MIN_REGIONS,
            "census_population_or_bootstrap_numeric_magnitudes_opened":0,
            "ecological_effects_computed":0,
            "scientific_decision":"HOLD_SOURCE_HEADERS_NOT_PROOF_OF_ISLAND_REPLICATION",
        }
    finally:
        wb.close()

def download():
    url=(f"https://{v28.HOST}/eds/api/dataset/{UUID}/object/download"
         f"?prefix={quote(OBJECT,safe='')}")
    with requests.Session() as session:
        raw,meta=v28.exact_get(session,url,v28.MAX_RAW)
    if len(raw)!=EXPECTED_BYTES or meta["source_bytes"]!=EXPECTED_BYTES:
        raise ValueError("SOURCE_STOP exact AADC original workbook byte size changed")
    return raw,meta

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--json",required=True,type=Path)
    args=ap.parse_args()
    receipt={"status":"HOLD_AADC_PUBLIC_ORIGINAL_SOURCE","ecological_effects_computed":0}
    try:
        raw,meta=download()
        structure=inspect(raw)
        receipt={
            "status":"V32_ORIGINAL_AADC_METAPOP_SOURCE_SCHEMA",
            "original_source_DOI":EXPECTED_SOURCE_DOI,
            "original_EDS_ID":"5516",
            "original_object_name":OBJECT,
            "source_SHA256":meta["sha256"],
            "original_file_bytes":len(raw),
            "source_verified_original_xlsx":True,
            "publisher_original_source_structure":structure,
            "bootstrap_draws_are_not_independent_islands":True,
            "record_level_individual_movement_observed":False,
            "scientific_decision":structure["scientific_decision"],
        }
    except Exception as exc:
        receipt["error_type"]=type(exc).__name__
        receipt["error"]=str(exc)[:600]
    args.json.parent.mkdir(parents=True,exist_ok=True)
    args.json.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
    if receipt["status"]=="HOLD_AADC_PUBLIC_ORIGINAL_SOURCE":
        raise SystemExit(2)

if __name__=="__main__":
    main()
