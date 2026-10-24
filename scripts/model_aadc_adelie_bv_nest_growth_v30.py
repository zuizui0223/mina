#!/usr/bin/env python3
"""V30: original AADC occupied-nest changes across actual Béchervaise / Verner islands."""
from __future__ import annotations
import argparse,io,json,math,re
from collections import Counter,defaultdict
from datetime import date,datetime
from pathlib import Path
from statistics import mean
from urllib.parse import quote
import requests
from openpyxl import load_workbook
import audit_aadc_three_island_original_v28 as v28
import audit_aadc_three_island_season_v29 as v29

PAIR=("Béchervaise","Verner")
FROZEN_YEARS=(1991,1996,1997,1998,1999,2000,2001,2002,2003,2004)
FROZEN_WINDOWS=tuple(range(1996,2004))
SHA=v29.FILES["three_island"]["sha"]

def original_bytes():
    spec=v29.FILES["three_island"]
    url=("https://data.aad.gov.au/eds/api/dataset/"+spec["uuid"]+
         "/object/download?prefix="+quote(spec["original"],safe=""))
    with requests.Session() as s:
        raw,meta=v28.exact_get(s,url,v28.MAX_RAW)
    if meta["sha256"]!=SHA:
        raise ValueError("SOURCE_STOP official AADC original 1103 workbook SHA drift")
    return raw,meta

def source_dates_and_field_support(sheet):
    header=list(next(sheet.iter_rows(min_row=2,max_row=2,values_only=True)))
    expected=("Season*","Béchervaise","Verner","Petersen")
    if tuple(str(x).strip() for x in header[:4])!=expected:
        raise ValueError("SCHEMA_STOP original publisher physical-island columns changed")
    byyear=defaultdict(list)
    rejected=Counter()
    for row in sheet.iter_rows(min_row=3,values_only=True):
        if not row:continue
        year=v29.source_year(row[0])
        if year is None or year<1990 or year>2004:
            continue
        stage=str(row[5] if len(row)>5 and row[5] is not None else "").strip().lower()
        fielddate=row[4] if len(row)>4 else None
        if stage!="occ nests":
            rejected["source_unit_not_literal_occ_nests"]+=1
            continue
        if isinstance(fielddate,datetime):
            fielddate=fielddate.date()
        if not isinstance(fielddate,date) or fielddate.year!=year or fielddate.month not in (11,12):
            rejected["source_date_not_verified_in_incubation_season"]+=1
            continue
        byyear[year].append(row)
    duplicates={year for year,rows in byyear.items() if len(rows)>1}
    unique={year:rows[0] for year,rows in byyear.items() if len(rows)==1}
    availability={label:{y for y,row in unique.items() if v29.count_present(row[pos])}
                  for pos,label in enumerate(("Béchervaise","Verner","Petersen"),start=1)}
    years=sorted(availability["Béchervaise"] & availability["Verner"])
    windows=sorted(y for y in years if y+1 in years)
    meta={
        "stage_source_literal":"occ nests",
        "primary_pair":list(PAIR),
        "source_stage_or_date_excluded_rows":dict(rejected),
        "duplicated_original_source_season_groups_excluded":len(duplicates),
        "source_count_available_years_by_named_island":{label:sorted(a)
                                                         for label,a in availability.items()},
        "stage_date_valid_original_years":sorted(unique),
        "original_Bechervaise_Verner_shared_years":years,
        "original_Bechervaise_Verner_adjacent_year_windows":windows,
        "stage_qualified_support_gate":(len(years)>=8 and len(windows)>=6),
        "exact_V29_outcome_blind_roster_reproduced":(
            years==list(FROZEN_YEARS) and windows==list(FROZEN_WINDOWS)),
        "source_count_magnitudes_opened":False,
    }
    return unique,years,windows,meta

def count_int(raw):
    if raw is None or isinstance(raw,bool):return None
    if isinstance(raw,(int,float)):
        x=float(raw)
    else:
        text=str(raw).strip().replace(",","")
        if not re.fullmatch(r"(?:0|[1-9]\d*)(?:\.0+)?",text):return None
        x=float(text)
    return int(x) if math.isfinite(x) and 0<=x<1e8 and x.is_integer() else None

def original_count_type(raw):
    """Classify publisher Excel cell *quality*, never export the count magnitude."""
    if raw is None or (isinstance(raw,str) and not raw.strip()):
        return "MISSING_CELL"
    if isinstance(raw,bool):
        return "BOOLEAN_NOT_COUNT"
    if isinstance(raw,(int,float)):
        n=float(raw)
        source="EXCEL_NUMERIC"
    else:
        source="TEXT"
        s=str(raw).strip().replace(",","")
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d+)?|\.\d+)(?:[eE][+-]?\d+)?",s):
            return "TEXT_NONNUMERIC"
        try:n=float(s)
        except ValueError:return "TEXT_NONNUMERIC"
    if not math.isfinite(n):
        return source+"_NONFINITE"
    if n<0:
        return source+"_NEGATIVE"
    if n>=1e8:
        return source+"_OUTSIDE_FROZEN_RANGE"
    if not n.is_integer():
        return source+"_FRACTIONAL"
    if source=="TEXT" and count_int(raw) is None:
        return "TEXT_INTEGER_NONCANONICAL"
    return source+"_NONNEGATIVE_INTEGER"

def pearson(xs,ys):
    if len(xs)!=len(ys) or len(xs)<3:return None
    xbar=mean(xs);ybar=mean(ys)
    a=[x-xbar for x in xs];b=[y-ybar for y in ys]
    den=math.sqrt(sum(x*x for x in a)*sum(y*y for y in b))
    return float(sum(x*y for x,y in zip(a,b))/den) if den>1e-12 else None

def real_growth(rows,years,windows):
    counts={}
    errors=Counter()
    original_cell_classes={}
    for year in years:
        row=rows[year]
        original_cell_classes[str(year)]={
            PAIR[0]:original_count_type(row[1]),
            PAIR[1]:original_count_type(row[2])
        }
        x=count_int(row[1]);y=count_int(row[2])
        if x is None or y is None:
            errors["invalid_numeric_original_occ_nests_year"]+=1
            continue
        counts[year]=(x,y)
    complete=[year for year in windows if year in counts and year+1 in counts]
    eligibility={
        "source_numeric_valid_pair_years":len(counts),
        "source_numeric_consecutive_windows":len(complete),
        "source_numeric_inclusion_errors":dict(errors),
        "source_publisher_cell_quality_by_year_island":original_cell_classes,
        "source_publisher_cell_quality_counts":dict(Counter(
            status for row in original_cell_classes.values() for status in row.values())),
        "numeric_cell_quality_diagnostic_does_not_display_counts":True,
        "original_numeric_years":sorted(counts),
        "original_adjacent_window_starts":complete
    }
    if len(counts)<8 or len(complete)<6 or complete!=list(FROZEN_WINDOWS):
        return {**eligibility,"status":"STOP_ORIGINAL_NUMERIC_SUPPORT_OR_YEAR_ROSTER",
                "ecological_effect_fitted":False}
    bg=[math.log1p(counts[y+1][0])-math.log1p(counts[y][0]) for y in complete]
    vg=[math.log1p(counts[y+1][1])-math.log1p(counts[y][1]) for y in complete]
    r=pearson(bg,vg)
    if r is None:
        return {**eligibility,"status":"STOP_ZERO_VARIANCE_ORIGINAL_PAIR",
                "ecological_effect_fitted":False}
    loo=[pearson(bg[:i]+bg[i+1:],vg[:i]+vg[i+1:]) for i in range(len(complete))]
    loo=[z for z in loo if z is not None]
    nonzero=sum(a!=0 and b!=0 for a,b in zip(bg,vg))
    samen=sum(a*b>0 for a,b in zip(bg,vg))
    cancelled=[]
    for y in complete:
        x=counts[y+1][0]-counts[y][0];z=counts[y+1][1]-counts[y][1]
        if abs(x)+abs(z)>0:
            cancelled.append(1-abs(x+z)/(abs(x)+abs(z)))
    return {**eligibility,
        "status":"V30_ORIGINAL_MAWSON_TWO_ISLAND_ADÉLIE_NEST_GROWTH",
        "ecological_effect_fitted":True,
        "all_eight_source_year_changes":{
            str(y):{"Bechervaise":a,"Verner":b} for y,a,b in zip(complete,bg,vg)},
        "nearby_physical_island_log1p_nest_growth_pearson_r":r,
        "leave_one_adjacent_year_out_r_min_max":[min(loo),max(loo)] if loo else None,
        "both_nonzero_growth_windows":nonzero,
        "same_direction_growth_windows":samen,
        "same_direction_share_nonzero":samen/nonzero if nonzero else None,
        "observed_nest_count_compensation_index_mean":mean(cancelled) if cancelled else None,
        "source_original_absolute_occupied_nest_counts_omitted":True,
        "two_islands_are_not_antactic_island_macro_replication":True,
        "individual_penguin_exchange_or_recruitment_inferred":False,
        "marine_climate_ice_effect_estimated":False,
        "no_inferential_p_values":True,
        "scientific_decision":"EXPLORATORY_SINGLE_EAST_ANTARCTIC_TWO_ISLAND_SYNCHRONY_ONLY",
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--json",required=True,type=Path)
    args=p.parse_args()
    receipt={"status":"HOLD_ORIGINAL_SOURCE","ecological_effect_fitted":False}
    try:
        raw,meta=original_bytes()
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            if v29.SHEET_MAIN not in wb.sheetnames:
                raise ValueError("SCHEMA_STOP original AADC source sheet vanished")
            good,years,windows,qa=source_dates_and_field_support(wb[v29.SHEET_MAIN])
            receipt={
                "status":"V30_ORIGINAL_AADC_TWO_ISLAND_STAGE_QA",
                "source_original_file_sha256":SHA,
                "source_original_bytes":meta["source_bytes"],
                "frozen_contract":"contracts/AADC_BV_ISLAND_NEST_GROWTH_V30.md",
                "original_stage_metadata":qa,
                "ecological_effect_fitted":False,
                "decision":"STOP_SOURCE_STAGE_SUPPORT_OR_YEAR_DRIFT"
            }
            if qa["stage_qualified_support_gate"] and qa["exact_V29_outcome_blind_roster_reproduced"]:
                fit=real_growth(good,years,windows)
                receipt["two_island_measured_response"]=fit
                receipt["ecological_effect_fitted"]=bool(fit["ecological_effect_fitted"])
                receipt["decision"]=fit.get("scientific_decision",fit["status"])
        finally:
            wb.close()
    except Exception as exc:
        receipt["error_type"]=type(exc).__name__
        receipt["error"]=str(exc)[:600]
    args.json.parent.mkdir(parents=True,exist_ok=True)
    args.json.write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
    if receipt["status"]=="HOLD_ORIGINAL_SOURCE":
        raise SystemExit(2)

if __name__=="__main__":
    main()
