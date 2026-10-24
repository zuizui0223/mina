#!/usr/bin/env python3
"""V31 authentic original AADC *continuous averaged occupied nest indices*.

Distinct measurement estimand from CLOSED V30 integer gate. Source methods
authorize non-integer reported nest estimates; no rounding or imputed zeros.
"""
from __future__ import annotations
import argparse
from collections import Counter
import io
import json
import math
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook
import model_aadc_adelie_bv_nest_growth_v30 as prior

FROZEN_YEARS=prior.FROZEN_YEARS
FROZEN_WINDOWS=prior.FROZEN_WINDOWS
SHA=prior.SHA
ORIGINAL_EXPECTED_FRACTIONAL=10
SOURCE_EXCEL_NUMERIC_TYPES=("EXCEL_NUMERIC_NONNEGATIVE_INTEGER","EXCEL_NUMERIC_FRACTIONAL")

def numeric_publisher_occupied_index(value):
    """Original publisher Excel VALUE; float is valid, no integer rounding."""
    if type(value) not in (int,float):
        return None
    value=float(value)
    if not math.isfinite(value) or value<0 or value>=1e8:
        return None
    return value

def summarize(good,years,windows):
    # This routine runs only after unchanged V29/V30 source/date/stage gate.
    typed={}
    reasons=Counter()
    value_classes={}
    for year in years:
        row=good[year]
        category={
            "Béchervaise":prior.original_count_type(row[1]),
            "Verner":prior.original_count_type(row[2]),
        }
        value_classes[str(year)]=category
        if any(kind not in SOURCE_EXCEL_NUMERIC_TYPES for kind in category.values()):
            reasons["NOT_FINITE_NONNEGATIVE_PUBLISHER_EXCEL_NUMERIC"]+=1
            continue
        b=numeric_publisher_occupied_index(row[1])
        v=numeric_publisher_occupied_index(row[2])
        if b is None or v is None:
            reasons["INVALID_PUBLISHER_NUMERIC_INDEX"]+=1
            continue
        typed[year]=(b,v)
    complete=[year for year in windows if year in typed and year+1 in typed]
    base={
        "status":"V31_PUBLISHER_CONTINUOUS_NEST_INDEX_SUPPORT",
        "source_census_years":sorted(typed),
        "original_count_value_types_by_year":value_classes,
        "original_numeric_quality_exclusions":dict(reasons),
        "source_numeric_years":len(typed),
        "source_numeric_consecutive_year_windows":len(complete),
        "original_adjacent_year_starts":complete,
        "prior_V30_integer_only_test_closed":True,
        "original_float_values_not_rounded":True,
        "original_source_count_values_not_reprinted":True,
        "ecological_effect_fitted":False,
    }
    if sorted(typed)!=list(FROZEN_YEARS) or complete!=list(FROZEN_WINDOWS):
        return {**base,"status":"STOP_PUBLISHER_NUMERIC_ORIGINAL_YEAR_ROSTER_DRIFT",
                "scientific_decision":"STOP_NO_CONTINUOUS_INDEX_SUPPORT"}
    bg=[math.log1p(typed[y+1][0])-math.log1p(typed[y][0]) for y in complete]
    vg=[math.log1p(typed[y+1][1])-math.log1p(typed[y][1]) for y in complete]
    r=prior.pearson(bg,vg)
    if r is None:
        return {**base,"status":"STOP_ORIGINAL_TWO_ISLAND_ZERO_VARIANCE",
                "scientific_decision":"STOP_ORIGINAL_TWO_ISLAND_ZERO_VARIANCE"}
    leave=[]
    for i in range(len(complete)):
        rr=prior.pearson(bg[:i]+bg[i+1:],vg[:i]+vg[i+1:])
        if rr is not None:leave.append(rr)
    signs_n=sum(x!=0 and z!=0 for x,z in zip(bg,vg))
    same=sum(x*z>0 for x,z in zip(bg,vg))
    arithmetic=[]
    for year in complete:
        db=typed[year+1][0]-typed[year][0]
        dv=typed[year+1][1]-typed[year][1]
        if abs(db)+abs(dv)>0:
            arithmetic.append(1-abs(db+dv)/(abs(db)+abs(dv)))
    return {
        **base,
        "status":"V31_ORIGINAL_AADC_TWO_PHYSICAL_ISLAND_REPORTED_NEST_INDEX_GROWTH",
        "ecological_effect_fitted":True,
        "source_Excel_numeric_fractional_year_island_count":sum(
            kind=="EXCEL_NUMERIC_FRACTIONAL"
            for one in value_classes.values() for kind in one.values()),
        "original_growth_start_years":list(complete),
        "island_year_source_derived_log1p_growth":{
            str(year):{"Bechervaise":float(b),"Verner":float(v)}
            for year,b,v in zip(complete,bg,vg)},
        "physical_island_pair_log1p_occupied_nest_index_growth_pearson_r":r,
        "leave_one_year_out_growth_r_min_max":[min(leave),max(leave)] if leave else None,
        "original_years_both_changes_nonzero":signs_n,
        "original_years_same_direction":same,
        "same_direction_fraction":same/signs_n if signs_n else None,
        "reported_index_count_change_arithmetic_cancellation_mean":mean(arithmetic) if arithmetic else None,
        "original_colony_level_counting_variance_unknown":True,
        "verner_specific_averaging_method_independently_verified":False,
        "individual_island_migration_measured":False,
        "marine_climate_causal_effect_identified":False,
        "independent_physical_islands":2,
        "no_inferential_p_value":True,
        "scientific_decision":"EXPLORATORY_SINGLE_NEARBY_ISLAND_NEST_INDEX_COVARIATION_NONCAUSAL"
    }

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--json",type=Path,required=True)
    a=parser.parse_args()
    report={"status":"HOLD_ORIGINAL_SOURCE","ecological_effect_fitted":False}
    try:
        raw,receipt=prior.original_bytes()
        if receipt["sha256"]!=SHA:
            raise ValueError("SOURCE_STOP original source SHA drift")
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        try:
            if prior.v29.SHEET_MAIN not in wb.sheetnames:
                raise ValueError("SCHEMA_STOP source sheet missing")
            good,years,windows,qa=prior.source_dates_and_field_support(wb[prior.v29.SHEET_MAIN])
            result={
                "status":"V31_ORIGINAL_SOURCE_STAGE_AND_MEASUREMENT_CHECK",
                "source_original_sha256":SHA,
                "source_original_bytes":receipt["source_bytes"],
                "source_accepted_count_unit":"author-reported possibly averaged occupied-nest index, floating values retained",
                "frozen_contract":"contracts/AADC_BV_CONTINUOUS_OCCUPIED_NEST_INDEX_V31.md",
                "source_stage_eligibility":qa,
                "v30_integer_support_result":"CLOSED_WITHOUT_EFFECT",
                "ecological_effect_fitted":False,
                "scientific_decision":"HOLD_ORIGINAL_SOURCE_STAGE"
            }
            if qa["stage_qualified_support_gate"] and qa["exact_V29_outcome_blind_roster_reproduced"]:
                analysis=summarize(good,years,windows)
                result["original_index_analysis"]=analysis
                result["ecological_effect_fitted"]=bool(analysis["ecological_effect_fitted"])
                result["scientific_decision"]=analysis["scientific_decision"]
            report=result
        finally:
            wb.close()
    except Exception as exc:
        report["error_type"]=type(exc).__name__
        report["error"]=str(exc)[:650]
    a.json.parent.mkdir(parents=True,exist_ok=True)
    a.json.write_text(json.dumps(report,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))
    if report["status"]=="HOLD_ORIGINAL_SOURCE":
        raise SystemExit(2)

if __name__=="__main__":
    main()
