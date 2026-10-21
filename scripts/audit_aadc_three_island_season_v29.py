#!/usr/bin/env python3
"""V29 original AADC three physical island matched year source-only support."""
from collections import defaultdict,Counter
import argparse
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import quote
import requests
from openpyxl import load_workbook
import io
import audit_aadc_three_island_original_v28 as v28

FILES={
 "three_island":{
     "uuid":"3509BE35-77D3-4179-BF21-09E981F512A1",
     "original":"population_counts_BI.xlsx",
     "sha":"8377fd17d4d84ccb791229ebdf5926605ae34aeffd0c1a21a0a7ae75d4facdf6"
 },
 "demography":{
     "uuid":"7673853e-9360-4563-a0e9-324300cfa0b7",
     "original":"Time series demography and environmental covariates.xlsx",
     "sha":"91862ed4f607c5031adbd220372a274b1376c68f535e5dcbc4b19412c1e48b67"
 }
}
ISLANDS=("Béchervaise","Verner","Petersen")
YEAR_MIN=1990
YEAR_MAX=2004
SHEET_MAIN="Population counts"
SHEET_OTHER="soe27"

def source_year(value):
    if value is None or isinstance(value,bool):
        return None
    if isinstance(value,(int,float)):
        if int(value)!=value:return None
        return int(value) if 1900<=value<=2050 else None
    text=str(value).strip()
    if re.fullmatch(r"(?:19|20)\d{2}",text):
        return int(text)
    match=re.fullmatch(r"((?:19|20)\d{2})\s*[-/]\s*(\d{2}|\d{4})",text)
    if not match:
        return None
    first=int(match.group(1))
    final=int(match.group(2))
    if final==(first+1)%100 or final==first+1:
        return first
    return None

def count_present(cell):
    if cell is None or isinstance(cell,bool):
        return False
    if isinstance(cell,(int,float)):
        return True
    s=str(cell).strip().replace(",","")
    return bool(re.fullmatch(r"(?:0|[1-9]\d*)(?:\.\d+)?",s))

def sheet_text_mask(ws,max_rows=50):
    """Only source text and non-identifying count PRESENCE; never count magnitude."""
    report=[]
    for row_no,row in enumerate(ws.iter_rows(values_only=True),start=1):
        if row_no>max_rows:break
        tags=[]
        for col_no,c in enumerate(row,1):
            if c is None or str(c).strip()=="":
                continue
            if isinstance(c,str):
                name=c.strip()
                if re.fullmatch(r"[+-]?\d+(?:\.\d+)?",name):
                    tag="NUM_PRESENT"
                else:
                    tag=name[:65]
            elif isinstance(c,(int,float)):
                tag="NUM_PRESENT"
            else:
                tag=str(type(c).__name__)
            # Original row season/year labels are source metadata, not outcome.
            if col_no==1 and source_year(c) is not None:
                tag="SEASON_START_"+str(source_year(c))
            tags.append([col_no,tag])
        if tags:
            report.append({"row":row_no,"nonempty_source_cell_types":tags})
    return report

def season_table(ws,startrow,minyear,maxyear,columns):
    """Generate immutable yearly presence/absence of ORIGINAL MEASUREMENT, not biological absence."""
    groups=defaultdict(list)
    bad=Counter()
    for index,row in enumerate(ws.iter_rows(values_only=True),start=1):
        if index<startrow or not row:
            continue
        yr=source_year(row[0])
        if yr is None:
            if row[0] is not None:bad["not_unambiguous_austral_season"]+=1
            continue
        if not minyear<=yr<=maxyear:continue
        groups[yr].append({label:count_present(row[i]) if i<len(row) else False
                           for label,i in columns.items()})
    duplicates={y for y,recs in groups.items() if len(recs)>1}
    eligible={y:rec[0] for y,rec in groups.items() if len(rec)==1}
    sets={label:{y for y,r in eligible.items() if r[label]} for label in columns}
    return sets,{
        "parsed_unique_season_labels_in_fixed_range":len(groups),
        "season_labels_with_duplicate_source_rows_excluded":sorted(duplicates),
        "source_years_with_single_original_row":len(eligible),
        "original_source_nonseason_row_labels_excluded":dict(bad),
        "lexically_numeric_reported_seasons_by_source_column":{name:sorted(yrs) for name,yrs in sets.items()},
        "is_original_numeric_cell_a_measured_zero":False,
        "original_bird_count_magnitudes_inspected":0
    }

def compare(sets):
    def cohort(names):
        years=sorted(set.intersection(*(sets[n] for n in names)))
        runs=sorted(y for y in years if y+1 in years)
        return {"original_source_island_name_roster":list(names),
                "common_original_seasons":len(years),
                "source_common_years":years,
                "strict_consecutive_annual_windows":len(runs),
                "strict_adjacent_window_start_years":runs,
                "frozen_next_stage_support_pass":len(years)>=8 and len(runs)>=6}
    return {
        "all_three":cohort(ISLANDS),
        "Bechervaise_Verner":cohort(ISLANDS[:2]),
        "Bechervaise_Petersen":cohort((ISLANDS[0],ISLANDS[2])),
        "Verner_Petersen":cohort(ISLANDS[1:]),
    }

def original_sources():
    original={}
    with requests.Session() as session:
        for label,spec in FILES.items():
            url=f"https://{v28.HOST}/eds/api/dataset/{quote(spec['uuid'])}/object/download?prefix={quote(spec['original'],safe='')}"
            content,metadata=v28.exact_get(session,url,v28.MAX_RAW)
            if metadata["sha256"]!=spec["sha"]:
                raise ValueError("SOURCE_STOP original AADC workbook version drift "+label)
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
            original[label]=book
    return original

def process(a,b):
    if SHEET_MAIN not in a.sheetnames or SHEET_OTHER not in a.sheetnames or "data" not in b.sheetnames:
        raise ValueError("SCHEMA_STOP original expected publisher worksheets missing")
    main=a[SHEET_MAIN]
    headers=list(next(main.iter_rows(min_row=2,max_row=2,values_only=True)))
    expected=("Season*","Béchervaise","Verner","Petersen")
    if tuple(str(x).strip() for x in headers[:4])!=expected:
        raise ValueError("SCHEMA_STOP three-island original publisher season/count columns changed")
    series,qa=season_table(main,3,YEAR_MIN,YEAR_MAX,{
        name:index+1 for index,name in enumerate(ISLANDS)
    })
    pairs=compare(series)
    soe=a[SHEET_OTHER]
    soe_header=[str(x).strip() if x is not None else "" for x in next(soe.iter_rows(values_only=True))]
    other_names={str(x).strip():i for i,x in enumerate(next(soe.iter_rows(values_only=True)))
                 if i>0 and isinstance(x,str) and x.strip()}
    islandnames=["Giganteus Island","Shirley Island"]
    sources_soe,otherQA=season_table(soe,2,1971,2005,{
        x:other_names[x] for x in islandnames if x in other_names
    })
    climate=b["data"]
    climate_header=list(next(climate.iter_rows(min_row=2,max_row=2,values_only=True)))
    yearcol=str(climate_header[0]).strip()
    if yearcol!="second.split.season":
        raise ValueError("SCHEMA_STOP environmental time series year column changed")
    environmental,met=season_table(climate,3,1990,2019,{
        "season_available":0,"site_air_temperature":6,"site_windchill":5,
        "nearshore_sea_ice":7
    })
    return {
        "status":"V29_ORIGINAL_AADC_TRUE_ISLAND_SEASON_SUPPORT_QA",
        "original_source_files_SHA256":{key:value["sha"] for key,value in FILES.items()},
        "frozen_fixed_island_season_range":[YEAR_MIN,YEAR_MAX],
        "original_main_sheet_header":expected,
        "main_source_original_metadata_QA":qa,
        "predeclared_three_island_pair_opportunities":pairs,
        "new_original_soe27_other_geographic_units":{
            "source_header":soe_header,
            "available_named_source_columns":list(other_names),
            "two_individual_named_island_year_presence":otherQA,
            "original_sheet_first_20_source_cell_type_masks":sheet_text_mask(soe,20),
            "note":"Shirley and Giganteus are different locality/programmes; source stage/identity not yet harmonized with Holme Bay and Rookery aggregates cannot be used as independent island replicates"
        },
        "independent_Bechervaise_environmental_time_series":{
            "source_headers_row_2":[str(x or "")[:90] for x in climate_header],
            "source_year_and_covariate_presence":met,
            "note":"Climate values are one geographic region near Béchervaise; not replicated for other distant islands"
        },
        "main_sheet_original_first_40_source_cell_type_masks":sheet_text_mask(main,40),
        "biological_nest_values_opened":0,
        "ecological_models_fitted":0,
        "scientific_decision":("PROCEED_STRICT_SOURCE_COUNT_UNIT_YEAR_VERIFICATION" if
           any(p["frozen_next_stage_support_pass"] for p in pairs.values())
           else "STOP_ORIGINAL_THREE_ISLAND_TEMPORAL_MATCH_SUPPORT")
    }

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--json",type=Path,required=True)
    args=ap.parse_args()
    receipt={"status":"HOLD_ORIGINAL_SOURCE","ecological_models_fitted":0}
    books={}
    try:
        books=original_sources()
        receipt=process(books["three_island"],books["demography"])
    except Exception as ex:
        receipt["error_type"]=type(ex).__name__
        receipt["error"]=str(ex)[:650]
    finally:
        for book in books.values():
            book.close()
    args.json.parent.mkdir(parents=True,exist_ok=True)
    args.json.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(receipt,ensure_ascii=False))
    if receipt["status"]=="HOLD_ORIGINAL_SOURCE":
        raise SystemExit(2)

if __name__=="__main__":
    main()
