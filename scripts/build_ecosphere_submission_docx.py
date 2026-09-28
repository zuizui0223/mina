#!/usr/bin/env python3
"""Build a structurally formatted Ecosphere Word Main Document preview.

The preview intentionally retains author-controlled placeholders. It is not a
submission-ready file until those placeholders are resolved and the rendered
document is visually inspected.
"""
from __future__ import annotations

import argparse
import copy
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from mina.ecosphere_submission_metadata import (
    affiliation_lines,
    ai_disclosure_suffix,
    author_line,
    corresponding_text,
    load_metadata,
    present_address_text,
    require_complete_metadata,
)

TITLE = (
    "Common decline, divergent endpoints: hierarchical demography across "
    "Antarctic penguin breeding islands"
)
MARKER = "[[SECTION_BREAK_AFTER_TITLE]]"

OPEN_RESEARCH = (
    "The primary Palmer Station Antarctica Long Term Ecological Research "
    "Adélie penguin census is publicly available at DOI "
    "10.6073/pasta/89dd52217ca37e3a72a67f7a9bc3c82e. The sea-ice and "
    "Palmer Station weather data used in the bounded mechanism tests are "
    "publicly available at DOI 10.6073/pasta/4207e529832840db2282498d9f4f4f05 "
    "and DOI 10.6073/pasta/3eefb45dbfb784c3cabe3690ea46fe9e, respectively; "
    "broader APBP/MAPPPD data are publicly available as cited in the manuscript. "
    "Novel analysis code, frozen endpoint contracts, and result receipts used "
    "for this submission are publicly accessible for peer review at "
    "https://github.com/zuizui0223/mina. If the manuscript is accepted, the "
    "exact version of record of the analysis code and derived outputs will be "
    "archived in a permanent repository with a DOI, and the final Open Research "
    "Statement will be updated with that identifier."
)

AI_DISCLOSURE = (
    "OpenAI ChatGPT (GPT-5.6 Sol) was used during analysis and manuscript "
    "development to assist with code drafting and review, statistical "
    "sensitivity-analysis scripting, literature searching, and editorial "
    "drafting. All analyses were executed from version-controlled code, "
    "numerical results were checked against frozen result receipts, cited "
    "literature was independently verified, and the authors remain responsible "
    "for all analyses, interpretations, and text."
)

METHODS_AI_DISCLOSURE = (
    "OpenAI ChatGPT (GPT-5.6 Sol) assisted with code drafting and review, "
    "sensitivity-analysis scripting, literature searching, and editorial "
    "drafting. Executed analyses and numerical outputs were checked against "
    "version-controlled code and frozen result receipts; the authors retain "
    "responsibility for all analytical and textual content."
)


UNDERBRACE_SOURCE = r"""\[
\beta_{total}
=\frac{\alpha_{sub}}{\gamma_{arch}}
=\underbrace{\frac{\alpha_{sub}}{\alpha_{island}}}_{\beta_{within}}
\underbrace{\frac{\alpha_{island}}{\gamma_{arch}}}_{\beta_{among}}.
\]"""

UNDERBRACE_REPLACEMENT = r"""\[
\beta_{total}
=\frac{\alpha_{sub}}{\gamma_{arch}}
=\beta_{within}\beta_{among}.
\]
\[
\beta_{within}=\frac{\alpha_{sub}}{\alpha_{island}},
\qquad
\beta_{among}=\frac{\alpha_{island}}{\gamma_{arch}}.
\]"""

OLD_DATA_AVAILABILITY = (
    "Analysis code, frozen endpoint contracts, result receipts and "
    "figure-building scripts will be supplied as an anonymized repository "
    "snapshot for peer review and archived with a permanent DOI on acceptance."
)

CURRENT_DATA_AVAILABILITY = (
    "Novel analysis code, frozen endpoint contracts, result receipts and "
    "figure-building scripts are publicly accessible for peer review at "
    "https://github.com/zuizui0223/mina. If accepted, the exact code and "
    "derived-output release will be archived in a permanent repository with "
    "a DOI."
)


def _normalize_submission_markdown(text: str) -> str:
    """Normalize Markdown constructs that do not round-trip cleanly to DOCX."""
    text = text.replace(
        "Following Wang and Loreau [@wang2014]",
        "Following @wang2014",
    )
    text = text.replace(UNDERBRACE_SOURCE, UNDERBRACE_REPLACEMENT)
    text = text.replace(OLD_DATA_AVAILABILITY, CURRENT_DATA_AVAILABILITY)

    # Pandoc/OMML can leave a stray glyph when an inline math expression ends
    # with '=' and bold Markdown begins immediately after the math delimiter,
    # e.g. \(\beta=\)**1.0111**. Keep the exact value inside the math span.
    text = re.sub(
        r"\\\(([^()\n]*?=)\\\)\*\*([^*\n]+)\*\*",
        lambda match: r"\(" + match.group(1) + match.group(2) + r"\)",
        text,
    )
    return text


def _strip_submission_source(text: str) -> str:
    lines = text.splitlines()
    out: list[str] = []
    skipping_reference_pointer = False
    for index, line in enumerate(lines):
        if index == 0 and line.startswith("# "):
            continue
        if line.startswith("**Ecosphere-oriented manuscript v0.6"):
            continue
        if line.startswith("**Keywords:**"):
            continue
        if line.strip() == "## References":
            skipping_reference_pointer = True
            continue
        if skipping_reference_pointer:
            continue
        out.append(line)

    cleaned = _normalize_submission_markdown("\n".join(out).rstrip())
    heading = "### Reproducibility and frozen result family"
    if heading not in cleaned:
        raise ValueError("reproducibility heading not found for AI disclosure")
    cleaned = cleaned.replace(
        heading,
        heading + "\n\n" + METHODS_AI_DISCLOSURE,
        1,
    )
    return cleaned


def _strip_caption_heading(text: str) -> str:
    lines = text.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    # Remove only the file-level H1 ("Figure captions v5 ..."). Preserve the
    # first actual caption heading ("## Figure 1 ...").
    if lines and lines[0].startswith("# ") and not lines[0].startswith("## "):
        lines.pop(0)
    while lines and not lines[0].strip():
        lines.pop(0)
    return _normalize_submission_markdown("\n".join(lines).strip())


def _open_research(metadata: dict[str, object] | None) -> str:
    if metadata is None:
        return OPEN_RESEARCH
    review_url=str(metadata.get("review_code_url") or "").strip()
    doi=str(metadata.get("permanent_archive_doi") or "").strip()
    text=OPEN_RESEARCH.replace(
        "https://github.com/zuizui0223/mina",
        review_url or "https://github.com/zuizui0223/mina",
    )
    if doi:
        old=(
            "If the manuscript is accepted, the exact version of record of "
            "the analysis code and derived outputs will be archived in a "
            "permanent repository with a DOI, and the final Open Research "
            "Statement will be updated with that identifier."
        )
        new=(
            "The exact version of record of the analysis code and derived "
            "outputs is archived at DOI " + doi + "."
        )
        text=text.replace(old,new)
    return text


def _metadata_blocks(
    metadata: dict[str, object] | None,
) -> dict[str, str]:
    if metadata is None:
        return {
            "authors":"[AUTHOR 1], [AUTHOR 2], [...]",
            "affiliations":"[AFFILIATIONS — REQUIRED]",
            "present":"[PRESENT ADDRESSES OR DELETE]",
            "corresponding":"[ONE AUTHOR NAME], [EMAIL ADDRESS]",
            "funding":"[FUNDING ACKNOWLEDGMENTS AND GRANT IDENTIFIERS — REQUIRED]",
            "additional_ack":"",
            "contributions":"[AUTHOR CONTRIBUTIONS — REQUIRED; confirm CRediT-style roles with every author.]",
            "coi":"[CONFLICT-OF-INTEREST STATEMENT — REQUIRED; confirm with every author.]",
            "ai_suffix":"",
        }

    present=present_address_text(metadata)
    additional=str(metadata.get("additional_acknowledgments") or "").strip()
    return {
        "authors":author_line(metadata),
        "affiliations":affiliation_lines(metadata),
        "present":present,
        "corresponding":corresponding_text(metadata),
        "funding":str(metadata.get("funding_acknowledgments") or "").strip(),
        "additional_ack":additional,
        "contributions":str(metadata.get("author_contributions") or "").strip(),
        "coi":str(metadata.get("conflict_of_interest") or "").strip(),
        "ai_suffix":ai_disclosure_suffix(metadata),
    }


def _combined_markdown(
    manuscript: str,
    captions: str,
    metadata: dict[str, object] | None = None,
) -> str:
    body = _strip_submission_source(manuscript)
    caption_body = _strip_caption_heading(captions)
    blocks=_metadata_blocks(metadata)
    open_research=_open_research(metadata)
    methods_disclosure=METHODS_AI_DISCLOSURE + blocks["ai_suffix"]
    body=body.replace(
        METHODS_AI_DISCLOSURE,
        methods_disclosure,
        1,
    )

    present_line=(
        f"**Present address(es), if applicable:** {blocks['present']}\n\n"
        if metadata is None or blocks["present"]!="None."
        else ""
    )
    title = f"""Ecosphere

Article — Animal Ecology

# {TITLE}

**Authors:** {blocks["authors"]}

**Affiliations:**  
{blocks["affiliations"]}

{present_line}**Corresponding author:** {blocks["corresponding"]}

## Open Research Statement

{open_research}

## Key words/phrases

Adélie penguin; breeding patches; hierarchical variability; island ecology; long-term monitoring; population dynamics; spatial synchrony; temporal compensation

{MARKER}

"""
    def sentence(value: str) -> str:
        value=value.strip()
        if not value:
            return ""
        return value if value[-1] in ".!?" else value + "."

    funding=sentence(blocks["funding"])
    additional=sentence(blocks["additional_ack"])
    acknowledgments_extra=" ".join(
        value for value in (funding,additional) if value
    )
    if acknowledgments_extra:
        acknowledgments_extra=" " + acknowledgments_extra

    backmatter = f"""

## Acknowledgments

We thank the Palmer Station Antarctica Long Term Ecological Research program and the field teams and data stewards who collected and curated the long-term penguin census at Palmer Station, Antarctica.{acknowledgments_extra} {AI_DISCLOSURE}{blocks["ai_suffix"]}

## Author Contributions

{blocks["contributions"]}

## Conflict of Interest Statement

{blocks["coi"]}

## References

::: {{#refs}}
:::

## Figure captions

{caption_body}
"""
    return title + body + backmatter


def _set_section_layout(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def _remove_line_numbering(sect_pr) -> None:
    for node in list(sect_pr.findall(qn("w:lnNumType"))):
        sect_pr.remove(node)


def _add_line_numbering(sect_pr) -> None:
    _remove_line_numbering(sect_pr)
    node = OxmlElement("w:lnNumType")
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:start"), "0")
    node.set(qn("w:restart"), "newSection")
    node.set(qn("w:distance"), "360")
    sect_pr.append(node)


def _suppress_line_number(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:suppressLineNumbers")) is None:
        p_pr.append(OxmlElement("w:suppressLineNumbers"))


def _page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _suppress_line_number(paragraph)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _split_title_section(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    marker_index = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip() == MARKER
        ),
        None,
    )
    if marker_index is None:
        raise ValueError("title/body section marker missing")

    marker = paragraphs[marker_index]
    boundary = next(
        (
            paragraph
            for paragraph in reversed(paragraphs[:marker_index])
            if paragraph.text.strip()
        ),
        None,
    )
    if boundary is None:
        raise ValueError("title-page boundary paragraph missing")

    body_sect_pr = doc.element.body.sectPr
    title_sect_pr = copy.deepcopy(body_sect_pr)
    _remove_line_numbering(title_sect_pr)

    type_node = title_sect_pr.find(qn("w:type"))
    if type_node is None:
        type_node = OxmlElement("w:type")
        title_sect_pr.insert(0, type_node)
    type_node.set(qn("w:val"), "nextPage")

    boundary_p_pr = boundary._p.get_or_add_pPr()
    existing = boundary_p_pr.find(qn("w:sectPr"))
    if existing is not None:
        boundary_p_pr.remove(existing)
    boundary_p_pr.append(title_sect_pr)
    _suppress_line_number(boundary)

    parent = marker._p.getparent()
    if parent is None:
        raise ValueError("section marker is detached")
    parent.remove(marker._p)

    _add_line_numbering(body_sect_pr)


def _apply_word_format(doc: Document) -> None:
    for section in doc.sections:
        _set_section_layout(section)

    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    def format_paragraph(paragraph, line_spacing: float) -> None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = line_spacing
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run._element.get_or_add_rPr().rFonts.set(
                qn("w:eastAsia"), "Times New Roman"
            )

    body_started = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Abstract":
            body_started = True
        format_paragraph(paragraph, 2.0 if body_started else 1.0)
        if not body_started:
            _suppress_line_number(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    format_paragraph(paragraph, 2.0)

    for index, section in enumerate(doc.sections):
        section.footer.is_linked_to_previous = False
        _page_number(section.footer.paragraphs[0])
        _remove_line_numbering(section._sectPr)
        if index > 0:
            _add_line_numbering(section._sectPr)

    props = doc.core_properties
    props.title = TITLE
    props.subject = "Ecosphere Article submission preview"
    props.author = ""
    props.last_modified_by = ""


def build(
    manuscript_path: Path,
    bibliography_path: Path,
    captions_path: Path,
    out_path: Path,
    metadata_path: Path | None = None,
    require_complete: bool = False,
) -> None:
    manuscript = manuscript_path.read_text(encoding="utf-8")
    captions = captions_path.read_text(encoding="utf-8")
    metadata=None
    if metadata_path is not None:
        metadata=load_metadata(metadata_path)
        if require_complete:
            require_complete_metadata(metadata)
    source = _combined_markdown(manuscript, captions, metadata=metadata)

    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        source_path = tmpdir / "submission.md"
        raw_docx = tmpdir / "submission_raw.docx"
        source_path.write_text(source, encoding="utf-8")
        subprocess.run(
            [
                "pandoc",
                str(source_path),
                "--from=markdown+raw_attribute+tex_math_single_backslash",
                "--to=docx",
                "--citeproc",
                f"--bibliography={bibliography_path}",
                "-o",
                str(raw_docx),
            ],
            check=True,
        )
        doc = Document(raw_docx)
        _split_title_section(doc)
        doc.save(tmpdir / "split.docx")
        doc = Document(tmpdir / "split.docx")
        _apply_word_format(doc)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manuscript", required=True, type=Path)
    parser.add_argument("--bibliography", required=True, type=Path)
    parser.add_argument("--captions", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--metadata", type=Path)
    parser.add_argument("--require-complete-metadata", action="store_true")
    args = parser.parse_args()
    build(
        args.manuscript,
        args.bibliography,
        args.captions,
        args.out,
        metadata_path=args.metadata,
        require_complete=args.require_complete_metadata,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
