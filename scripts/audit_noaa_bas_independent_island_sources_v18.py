#!/usr/bin/env python3
"""V18 source-only new NOAA/BAS original island breeding time series audit."""
import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import re
from urllib.parse import urlparse
import requests
from openpyxl import load_workbook

NOAA_URL="https://www.ncei.noaa.gov/data/oceans/archive/arc0132/0185113/1.1/data/0-data/Annual%20Penguin%20Census%201977-2015%20v1_0.xlsx"
BAS_URL="https://ramadda.data.bas.ac.uk/repository/entry/get/signy_gentoo_breeding_success.csv?entryid=synth%3A883c6b40-024e-41e8-9c3e-2668a44b5751%3AL3NpZ255X2dlbnRvb19icmVlZGluZ19zdWNjZXNzLmNzdg%3D%3D"
EXPECTED_NOAA={"Copa_Adelie","Copa_Gentoo","Copa_Chinstrap",
               "CapeShirreff_Chinstrap","CapeShirreff_Gentoo"}
LIMIT_BYTES=2*1024*1024
USER_AGENT="mina-antarctic-public-original-time-series-schema-audit/1.0"

def source_get(url,allowed_host):
    if urlparse(url).scheme!="https":
        raise ValueError("SOURCE_STOP only verified official HTTPS")
    with requests.Session() as session:
        r=session.get(url,headers={"User-Agent":USER_AGENT},timeout=30,
                      stream=True,allow_redirects=True)
        final=urlparse(r.url)
        if final.scheme!="https":
            raise ValueError("SOURCE_STOP insecure redirect")
        hostname=(final.hostname or "").lower()
        if allowed_host=="noaa":
            ok=hostname=="www.ncei.noaa.gov" or hostname.endswith(".noaa.gov")
        else:
            ok=hostname=="ramadda.data.bas.ac.uk" or hostname.endswith(".bas.ac.uk")
        if not ok:
            raise ValueError("SOURCE_STOP redirected outside official original host")
        r.raise_for_status()
        parts=[]
        size=0
        for block in r.iter_content(chunk_size=16384):
            size+=len(block)
            if size>LIMIT_BYTES:
                raise ValueError("SOURCE_STOP original source exceeds bounded size")
            parts.append(block)
        data=b"".join(parts)
        if not data:
            raise ValueError("SOURCE_STOP no source body")
        return data,dict(final_url=r.url,mime=r.headers.get("Content-Type"),
                         bytes=len(data),sha256=hashlib.sha256(data).hexdigest())

def normal_header(value):
    return re.sub(r"[^a-z0-9]","",str(value).lower()) if value is not None else ""

def check_noaa(raw):
    if not raw.startswith(b"PK"):
        raise ValueError("SOURCE_STOP NOAA returned non-XLSX original")
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    report={"publisher_workbook_sheets":list(wb.sheetnames)}
    try:
        sheets=[]
        for ws in wb:
            top=[]
            for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
                if i>12:break
                # only source text headers; never display Count magnitudes.
                labels=[str(v).strip() if isinstance(v,str) else ""
                        for v in row]
                top.append((i,labels))
            target={normal_header(x):x for x in EXPECTED_NOAA}
            selected=None
            for n,names in top:
                hits={h:names.index(h) for h in names if normal_header(h) in target}
                actual={normal_header(h):i for i,h in enumerate(names) if normal_header(h) in target}
                if len(actual)>=3:
                    selected=(n,names,actual)
                    break
            row_summary={"name":ws.title,"declared_max_row":ws.max_row,
                         "declared_max_column":ws.max_column,
                         "source_top_header_text_candidates":[
                             {"row":n,"text_headers":[v for v in values if v][:16]}
                             for n,values in top if sum(bool(v) for v in values)>1][:5],
                         "publisher_five_series_headers_found":0}
            if selected:
                header_row,names,actual=selected
                counts={target[key]:0 for key in actual}
                original_year_labels=0
                # Metadata-only: only presence/absence of cell, no numeric magnitude.
                for row in ws.iter_rows(min_row=header_row+1,values_only=True):
                    if any(v is not None and str(v).strip() for v in row):
                        if len(row)>0 and row[0] is not None and str(row[0]).strip():
                            original_year_labels+=1
                        for key,col in actual.items():
                            if col<len(row) and row[col] is not None and str(row[col]).strip():
                                counts[target[key]]+=1
                row_summary.update({
                    "source_header_row":header_row,
                    "publisher_five_series_headers_found":len(actual),
                    "original_nonblank_by_publisher_series":counts,
                    "nonblank_austral_year_labels":original_year_labels,
                    "original_reported_columns":names[:25]
                })
            sheets.append(row_summary)
        report["original_sheets_structure"]=sheets
        report["any_sheet_contains_all_publisher_five_series"]=any(
            x["publisher_five_series_headers_found"]==5 for x in sheets)
        return report
    finally:
        wb.close()

def check_bas(raw):
    if raw[:80].lower().lstrip().startswith((b"<html",b"<!doctype",b"<script")):
        raise ValueError("SOURCE_STOP BAS response is HTML not original CSV")
    try:
        text=raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text=raw.decode("latin-1")
    reader=csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("SOURCE_STOP BAS original CSV missing headings")
    header=[str(x).strip() for x in reader.fieldnames if x is not None]
    if len(header)<3 or len(header)!=len(set(header)):
        raise ValueError("SCHEMA_STOP BAS original columns missing or repeated")
    rows=0
    nonblank={name:0 for name in header}
    for row in reader:
        rows+=1
        # Structural completeness ONLY. Values are not recorded/used.
        for name in header:
            if row.get(name) is not None and str(row[name]).strip():
                nonblank[name]+=1
    return {"original_declared_headers":header,
            "original_data_rows":rows,
            "per_column_nonblank_cells":nonblank,
            "numeric_stage_value_cells_opened":0}

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--out",type=Path,required=True)
    args=p.parse_args()
    report={"status":"V18_SOURCE_STAGE_CHECKED_NO_BIOLOGY",
            "scientific_decision":"HOLD_SOURCE_COMPARABILITY_AND_INDEPENDENT_ENVIRONMENT",
            "noaa_DOI":"10.25921/58yq-7g68",
            "bas_gentoo_DOI":"10.5285/883c6b40-024e-41e8-9c3e-2668a44b5751",
            "biological_count_magnitudes_opened":0,
            "biological_effect_models_fit":0,
            "noaabas_count_unit_equality_assumed":False,
            "source_records_are_not_independent_islands":True}
    for label,url,allowed,parser in (
        ("noaa",NOAA_URL,"noaa",check_noaa),
        ("bas_gentoo",BAS_URL,"bas",check_bas)):
        try:
            raw,metadata=source_get(url,allowed)
            report[label]={"download_status":"ORIGINAL_SOURCE_RETRIEVED",
                           "source_receipt":metadata,"structure":parser(raw)}
        except Exception as exc:
            report[label]={"download_status":"HOLD_ORIGINAL_SOURCE_OR_SCHEMA",
                           "error_type":type(exc).__name__,"error":str(exc)[:550]}
    if report.get("noaa",{}).get("download_status")!="ORIGINAL_SOURCE_RETRIEVED":
        report["scientific_decision"]="HOLD_NOAA_ORIGINAL_DOWNLOAD"
    elif not report["noaa"]["structure"]["any_sheet_contains_all_publisher_five_series"]:
        report["scientific_decision"]="HOLD_NOAA_FIVE_SERIES_HEADER_NOT_VERIFIED"
    elif report.get("bas_gentoo",{}).get("download_status")!="ORIGINAL_SOURCE_RETRIEVED":
        report["scientific_decision"]="HOLD_BAS_ORIGINAL_GENTOO_ACCESS"
    else:
        report["scientific_decision"]="PASS_ORIGINAL_TWO_ARCHIVES_NOT_EQUIVALENT_METRICS_NO_MODEL"
    args.out.parent.mkdir(parents=True,exist_ok=True)
    args.out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))
    if report["scientific_decision"]=="HOLD_NOAA_ORIGINAL_DOWNLOAD":
        raise SystemExit(2)

if __name__=="__main__":
    main()
